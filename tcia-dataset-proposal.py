import streamlit as st
import pandas as pd
import os
import io
import json
import zipfile
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
from docx import Document
import PyPDF2
import datetime
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
import re
from xml.sax.saxutils import escape

from tcia_proposals.proposal_metadata import (
    DATA_MODELS_URL, label_options, validate_components, validate_resources,
    component_text, resource_text, add_curation_placeholders,
    category_options, data_type_options, data_type_display, normalize_resource_url, load_form_reference, species_options,
)
from tcia_proposals.pdf_footer import draw_dsa_footer, replace_dsa_page_number

st.set_page_config(page_title="TCIA Dataset Proposal Form", layout="wide")

# Constants & Configuration
RESOURCES_DIR = os.path.join(os.path.dirname(__file__), 'tcia_proposals', 'resources')

LABELS = {
    "Proposal Type": "What kind of dataset are you submitting?",
    "Scientific POC Name": "Scientific POC Name*",
    "Scientific POC Email": "Scientific POC Institutional Email*",
    "Scientific POC Phone": "Scientific POC Phone",
    "Technical POC Name": "Technical POC Name*",
    "Technical POC Email": "Technical POC Institutional Email*",
    "Technical POC Phone": "Technical POC Phone",
    "Legal POC Name": "Legal POC Name*",
    "Legal POC Email": "Legal POC Institutional Email*",
    "Legal POC Phone": "Legal POC Phone",
    "Title": "Suggest a descriptive title for your dataset*",
    "Nickname": "Suggest a shorter nickname for your dataset*",
    "Authors": "List the authors of this data set*",
    "Abstract": "Dataset Abstract*",
    "Published Elsewhere": "Has this data ever been published elsewhere?*",
    "disease_site": "Primary disease site/location*",
    "diagnosis": "Disease type or cohort context*",
    "image_types": "Which image types are included in the data set?*",
    "supporting_data": "Which kinds of supporting data are included in the data set?*",
    "file_formats": "Specify the file format utilized for each type of data*",
    "modifications": "Describe any steps taken to de-identify or otherwise modify the data in preparation for submission*",
    "faces": "Does your data contain any images of patient faces?*",
    "collections_analyzed": "Which TCIA collection(s) did you analyze?*",
    "derived_types": "What types of derived data are included in the dataset?*",
    "disk_space": "Approximate disk space required*",
    "Time Constraints": "Are there any time constraints associated with sharing your data set?*",
    "manuscripts": "Are there any manuscripts or preprints that might help us better understand your dataset?",
    "adult_or_childhood_study": "Is this an Adult or Childhood study?*",
    "number_of_subjects": "Approximate number of subjects*",
    "acknowledgements": "Acknowledgements*",
    "funding_agency": "Funding Agency",
    "funding_source_program_name": "Funding Source Program Name",
    "grant_id": "Grant ID",
    "why_tcia": "Why would you like to publish this dataset on TCIA?*",
    "software_code": "Do you have any related resources such as source code, Jupyter notebooks, web sites or other software that will help users work with your data?*",
    "software_details": "Details*",
    "species": "Species or subject type*",
    "data_components": "Data being submitted to TCIA",
    "external_resources": "External resources",
    "label_catalog_version": "Data model reference version",
    "proposal_model_version": "Proposal selection guidance version"
}

@st.cache_data(ttl=3600, show_spinner=False)
def current_form_reference():
    return load_form_reference()


reference = current_form_reference()
catalog = reference['catalog']
proposal_model = reference['proposal_model']


def model_link(field, text="Learn about these terms"):
    st.markdown(f"[{text}]({DATA_MODELS_URL}#{field.replace('_', '-')})")


def choose_labels(field, label, key):
    options = label_options(catalog, field)
    if "Other" not in options:
        options.append("Other")
    with st.container(border=True):
        selected = st.multiselect(label, options=options, key=key,
                                 help="Choose Other if you don't see a good match for your dataset.")
        if "Other" in selected:
            details = st.text_input("Please describe the other disease or site", key=f"{key}_other")
            if not details.strip():
                input_errors.append(f"{label}: describe your Other selection.")
            result = [value if value != "Other" else f"Other - {details.strip()}" for value in selected]
        else:
            result = selected
        model_link(field)
    return result


def row_ids(name, initial=False):
    if name not in st.session_state:
        st.session_state[name] = [0] if initial else []
        st.session_state[f"{name}_next"] = 1 if initial else 0
    return st.session_state[name]


def add_row(name):
    next_id = st.session_state[f"{name}_next"]
    st.session_state[name].append(next_id)
    st.session_state[f"{name}_next"] = next_id + 1


def remove_row(name, row_id):
    st.session_state[name].remove(row_id)


def single_choice(label, field, key):
    options = category_options(catalog, proposal_model) if field == 'data_category' else label_options(catalog, field)
    if field == 'file_type':
        options = [value for value in options if value != 'ZIP']
    if "Other" not in options:
        options.append("Other")
    return st.selectbox(label, options, index=None, placeholder="Select an option", key=key)


def add_manuscript(source):
    category = st.session_state.ms_cat_new
    if source == 'URL':
        value = st.session_state.get('ms_url_new', '').strip()
        if not value:
            st.session_state.manuscript_error = 'Please provide a URL.'
            return
        st.session_state.manuscript_list.append({'category': category, 'type': 'URL', 'url': value})
        st.session_state.ms_url_new = ''
    else:
        key = f"ms_file_new_{st.session_state.get('ms_file_version', 0)}"
        value = st.session_state.get(key)
        if not value:
            st.session_state.manuscript_error = 'Please upload a file.'
            return
        st.session_state.manuscript_list.append({'category': category, 'type': 'File', 'file': value})
        st.session_state.ms_file_version = st.session_state.get('ms_file_version', 0) + 1
    st.session_state.manuscript_error = ''


def remove_manuscript(index):
    st.session_state.manuscript_list.pop(index)


def normalize_url_field(key):
    st.session_state[key] = normalize_resource_url(st.session_state.get(key, ''))


def yes_no_details(label, key, detail_label, help_text=None):
    answer = st.radio(label, ['No', 'Yes'], index=0, key=key, horizontal=True)
    if answer == 'Yes':
        details = st.text_area(detail_label, key=f'{key}_details', help=help_text)
        if not details.strip():
            input_errors.append(f'{label}: provide details for your Yes answer.')
        return f'Yes - {details.strip()}'
    return 'No'


AGREEMENT_TEMPLATE = os.path.join(RESOURCES_DIR, 'agreement_template.pdf')
HELP_DESK_EMAIL = os.getenv("TCIA_HELP_DESK_EMAIL", "help@cancerimagingarchive.net")
INSTITUTIONAL_EMAIL_HELP = (
    "Please use an email address provided by your institution rather than a "
    "free personal email service such as Gmail."
)

# Server-side SMTP configuration
SMTP_SERVER = os.getenv("SMTP_SERVER")
SMTP_PORT = int(os.getenv("SMTP_PORT", 587))
SMTP_USER = os.getenv("SMTP_USER")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD")

st.markdown(
    """
    <div style="display: flex; align-items: center; margin-bottom: 20px;">
        <img src="https://www.cancerimagingarchive.net/wp-content/uploads/2021/06/TCIA-Logo-01.png" style="height: 75px; margin-right: 15px;">
        <h1 style="margin: 0;">Dataset Proposal Form</h1>
    </div>
    """,
    unsafe_allow_html=True
)


# Initial choice
proposal_type = st.radio(
    "What kind of dataset are you submitting?",
    options=["New Collection Proposal", "Analysis Results Proposal"],
    help="Select 'New Collection' if you are submitting primary imaging data. Select 'Analysis Results' if you are submitting derived data (e.g., segmentations) from existing TCIA collections.",
    key="prop_type"
)

if proposal_type == "New Collection Proposal":
    st.info("💡 **New Collection Proposal**: Used for contributing original imaging data that has not been previously hosted on TCIA.")
else:
    st.info("💡 **Analysis Results Proposal**: Used for contributing derived data like annotations, segmentations, or radiomics features based on existing TCIA datasets.")

st.markdown("---")

# Form Questions
st.subheader("Contact Information")
if SMTP_SERVER and SMTP_USER and SMTP_PASSWORD:
    st.info("💡 **Note**: After completing this form, a copy of this proposal will be emailed to the Scientific, Technical, and Legal points of contact listed below.")

col1, col2, col3 = st.columns(3)
with col1:
    sci_poc_name = st.text_input(LABELS["Scientific POC Name"], help="The person to contact about the proposal and data collection.", key="sci_poc_name")
    tech_poc_name = st.text_input(LABELS["Technical POC Name"], help="The person involved in sending the data.", key="tech_poc_name")
    legal_poc_name = st.text_input(LABELS["Legal POC Name"], help="Authorized signatory who will sign the TCIA Data Submission Agreement. This should not be the PI or department chair.", key="legal_poc_name")
with col2:
    sci_poc_email = st.text_input(LABELS["Scientific POC Email"], help=INSTITUTIONAL_EMAIL_HELP, key="sci_poc_email")
    tech_poc_email = st.text_input(LABELS["Technical POC Email"], help=INSTITUTIONAL_EMAIL_HELP, key="tech_poc_email")
    legal_poc_email = st.text_input(LABELS["Legal POC Email"], help=INSTITUTIONAL_EMAIL_HELP, key="legal_poc_email")
with col3:
    sci_poc_phone = st.text_input(LABELS["Scientific POC Phone"], key="sci_poc_phone", placeholder="+1-555-555-5555")
    tech_poc_phone = st.text_input(LABELS["Technical POC Phone"], key="tech_poc_phone", placeholder="+1-555-555-5555")
    legal_poc_phone = st.text_input(LABELS["Legal POC Phone"], key="legal_poc_phone", placeholder="+1-555-555-5555")

st.subheader("Dataset Publication Details")
title = st.text_input(LABELS["Title"], help="Similar to a manuscript title.", key="title")
nickname = st.text_input(LABELS["Nickname"], help="Must be < 30 characters, alphanumeric and dashes only.", max_chars=30, key="nickname")
if nickname and not re.match(r"^[a-zA-Z0-9-]+$", nickname):
    st.error("⚠️ Invalid nickname. Only letters, numbers, and dashes are allowed.")

# Authors Section
st.write(f"**{LABELS['Authors']}**")
st.markdown("""
Please list all authors in the order you'd like them to appear if the data set is published on TCIA.
Names should be listed as **(FAMILY, GIVEN) - ORCID**.
Please use **semicolons** to separate authors, or list them **one per line**.
You can quickly register for an ORCID at [https://orcid.org/signin](https://orcid.org/signin) if you don't have one already.
""")
authors_raw = st.text_area("Example: Smith, John - 0000-0002-1234-5678; Doe, Jane",
                           help="List authors one per line or separated by semicolons.",
                           key="authors_raw_input")

st.write(f"**{LABELS['Abstract']}**")
st.markdown("""
Provide a brief overview of the dataset under 1000 characters, including:
1. Number of subjects
2. Types of imaging data included
3. Types of non-imaging supporting data (e.g., image classifications, segmentations, demographics, treatment details, outcomes)
4. Potential research applications of the dataset.
""")
abstract = st.text_area(LABELS["Abstract"], label_visibility="collapsed", help="Focus on describing the dataset itself.", max_chars=1000, key="abstract")

# Manuscripts Section
st.write(f"**{LABELS['manuscripts']}**")
st.markdown("""
We highly encourage "data descriptor" preprints published on [arxiv.org](https://arxiv.org/) that are focused on the contents of the dataset (not scientific advances) with [CICADAS](https://cancerimagingarchive.net/cicadas) formatting preferred.
""")
has_manuscripts = st.radio(LABELS["manuscripts"], options=["No", "Yes"], index=0, key="has_manuscripts", label_visibility="collapsed")

if 'manuscript_list' not in st.session_state:
    st.session_state.manuscript_list = []

if has_manuscripts == "Yes":
    # Display existing manuscripts
    for idx, ms in enumerate(st.session_state.manuscript_list):
        with st.container(border=True):
            col_info, col_del = st.columns([6, 1])
            with col_info:
                source_val = ms['url'] if ms['type'] == 'URL' else ms['file'].name
                st.write(f"**{ms['category']}**: {source_val}")
            with col_del:
                st.button("🗑️", key=f"del_ms_{idx}", on_click=remove_manuscript, args=(idx,))

    # Add new manuscript
    with st.expander("➕ Add a Manuscript/Preprint"):
        ms_category = st.selectbox("Category", options=["Dataset Descriptor", "Regular manuscript"], key="ms_cat_new")
        ms_type = st.radio("Source", options=["URL", "Upload File"], key="ms_type_new")

        if ms_type == "URL":
            ms_url = st.text_input("Enter URL", key="ms_url_new", placeholder="https://doi.org/10.1101/...")
            st.button("Add URL to List", on_click=add_manuscript, args=('URL',))
        else:
            file_key = f"ms_file_new_{st.session_state.get('ms_file_version', 0)}"
            st.file_uploader("Upload PDF or Word document", type=["pdf", "docx", "doc"], key=file_key)
            st.button("Add File to List", on_click=add_manuscript, args=('File',))
        if st.session_state.get('manuscript_error'):
            st.error(st.session_state.manuscript_error)



st.subheader("Data Collection Details")

# Dataset-level labels are separate from component and external-resource labels.
extra_data = {}
input_errors = []
if proposal_type == "Analysis Results Proposal":
    extra_data['collections_analyzed'] = st.text_input(LABELS["collections_analyzed"], key="collections_analyzed")
    st.caption("Describe the subjects and disease context represented in your submitted analysis subset below.")

col1, col2 = st.columns(2)
with col1:
    extra_data['disease_site'] = choose_labels('disease_location', LABELS['disease_site'], 'disease_site_ms')
with col2:
    extra_data['diagnosis'] = choose_labels('disease_type', LABELS['diagnosis'], 'diagnosis_ms')

extra_data['species'] = st.multiselect(LABELS['species'], species_options(catalog, proposal_model), default=['Human'], key='species')
model_link('species')
if 'Human' in extra_data['species']:
    extra_data['adult_or_childhood_study'] = st.multiselect(
        LABELS['adult_or_childhood_study'], ['Adolescent and Young Adult', 'Adult', 'Pediatric'],
        default=['Adult'], key='adult_or_childhood_study')

st.subheader("Data being submitted to TCIA")
col1, col2 = st.columns(2)
with col1:
    extra_data['number_of_subjects'] = st.text_input(LABELS["number_of_subjects"], key="number_of_subjects")
with col2:
    extra_data['disk_space'] = st.text_input(LABELS["disk_space"], key="disk_space")

st.markdown("Add a row for each group of files with **one category and one file format**. "
            "You may select multiple data types in a row. Use separate rows for different categories or formats.")
st.caption("Examples: Radiology Images / CT and MR / DICOM; Clinical Data / Follow-Up, Demographic and Treatment / CSV.")
st.markdown(f"Learn about these terms: [Categories]({DATA_MODELS_URL}#data-category) · "
            f"[Data types]({DATA_MODELS_URL}#data-type) · [File formats]({DATA_MODELS_URL}#file-type)")
components = []
for number, row_id in enumerate(row_ids('component_rows', initial=True), 1):
    with st.container(border=True):
        st.write(f"**Data component {number}**")
        category = single_choice("Data category", 'data_category', f'component_{row_id}_category')
        type_options = data_type_options(catalog, category, proposal_model)
        type_key = f'component_{row_id}_types'
        previous_types = st.session_state.get(type_key, [])
        kept_types = [value for value in previous_types if value in type_options]
        if kept_types != previous_types:
            st.session_state[type_key] = kept_types
            st.info("Some data types were cleared because they do not apply to the selected category.")
        types = st.multiselect("Data types", type_options, key=type_key,
                               format_func=lambda value: data_type_display(value, proposal_model), disabled=category is None,
                               help="Choose a category first. Select Other if its listed data types do not fit.")
        file_format = single_choice("File format", 'file_type', f'component_{row_id}_format')
        component = {'category': category, 'data_types': types, 'file_format': file_format}
        for field, selected in [('category', category), ('data_types', types), ('file_format', file_format)]:
            if (isinstance(selected, list) and 'Other' in selected) or selected == 'Other':
                component[f'other_{field}'] = st.text_input(
                    f"Describe the other {field.replace('_', ' ')}", key=f'component_{row_id}_other_{field}')
        components.append(component)
        st.button("Remove component", key=f'remove_component_{row_id}',
                  on_click=remove_row, args=('component_rows', row_id))
st.button("Add data component", on_click=add_row, args=('component_rows',))

extra_data['data_components'] = json.dumps(components)
extra_data['label_catalog_version'] = catalog['generated_at_utc']
extra_data['proposal_model_version'] = proposal_model['version']

if proposal_type == "New Collection Proposal":
    extra_data['modifications'] = st.text_area(LABELS["modifications"], key="modifications")
    extra_data['faces'] = st.radio(LABELS["faces"], options=["No", "Yes"], index=0, key="faces")
    if extra_data['faces'] == 'Yes':
        st.warning("Please review the [NIH Controlled Data Access Policy]"
                   "(https://www.cancerimagingarchive.net/nih-controlled-data-access-policy/) "
                   "for information about sharing images containing patient faces.")

st.subheader("External resources")
st.markdown("Optionally link related data or software hosted elsewhere. "
            "Add one category, URL, and resource context per resource. Include files being submitted to TCIA in the data components above.")
model_link('external_resource', 'Learn about external resource categories')
resources = []
for number, row_id in enumerate(row_ids('resource_rows'), 1):
    with st.container(border=True):
        st.write(f"**External resource {number}**")
        category = single_choice('Resource category', 'external_resource', f'resource_{row_id}_category')
        url_key = f'resource_{row_id}_url'
        url = st.text_input('Resource URL', key=url_key, placeholder='https://…',
                            help="You can enter a web address such as gdc.nci.nih.gov; https:// is added automatically.",
                            on_change=normalize_url_field, args=(url_key,))
        context = st.text_input('Resource Context', key=f'resource_{row_id}_context',
                                help='Explain what this resource includes.')
        resource = {'category': category, 'url': normalize_resource_url(url), 'context': context.strip()}
        if category == 'Other':
            resource['other_category'] = st.text_input('Describe the other resource category', key=f'resource_{row_id}_other')
        resources.append(resource)
        st.button('Remove resource', key=f'remove_resource_{row_id}',
                  on_click=remove_row, args=('resource_rows', row_id))
st.button('Add external resource', on_click=add_row, args=('resource_rows',))
extra_data['external_resources'] = json.dumps(resources)
# Shared bottom fields
col1, col2 = st.columns(2)
with col1:
    published_elsewhere = yes_no_details(
        LABELS["Published Elsewhere"], 'published_elsewhere', 'Please describe the previous publication',
        "Where was the data published? Why publish on TCIA, and will the original remain accessible?")
with col2:
    time_constraints = yes_no_details(
        LABELS["Time Constraints"], 'time_constraints', 'Please describe the time constraints')

extra_data['why_tcia'] = st.multiselect(
    LABELS["why_tcia"],
    options=["To meet a funding agency's requirements", "To meet a journal's requirements", "To facilitate collaboration", "To facilitate a challenge competition", "Other"],
    key="why_tcia"
)

button_label = "Submit" if (SMTP_SERVER and SMTP_USER and SMTP_PASSWORD) else "Generate Proposal Documents"
submit_button = st.button(button_label, type="primary")

# Processing after submission
if submit_button:
    # 1. Validation
    missing_fields = list(input_errors) + validate_components(components, proposal_model) + validate_resources(resources)
    if not sci_poc_name: missing_fields.append(LABELS["Scientific POC Name"])
    if not sci_poc_email: missing_fields.append(LABELS["Scientific POC Email"])
    if not tech_poc_name: missing_fields.append(LABELS["Technical POC Name"])
    if not tech_poc_email: missing_fields.append(LABELS["Technical POC Email"])
    if not legal_poc_name: missing_fields.append(LABELS["Legal POC Name"])
    if not legal_poc_email: missing_fields.append(LABELS["Legal POC Email"])
    if not title: missing_fields.append(LABELS["Title"])
    if not nickname:
        missing_fields.append(LABELS["Nickname"])
    elif not re.match(r"^[a-zA-Z0-9-]+$", nickname):
        st.error("⚠️ Invalid nickname. Please use only letters, numbers, and dashes.")
        st.stop()
    if not authors_raw:
        missing_fields.append(LABELS["Authors"])
    if not abstract: missing_fields.append(LABELS["Abstract"])
    if not published_elsewhere: missing_fields.append(LABELS["Published Elsewhere"])

    # Define optional fields that shouldn't block submission
    optional_fields = ['external_resources']

    for key, val in extra_data.items():
        if key in optional_fields:
            continue
        if not val:
            label = LABELS.get(key, key.replace('_', ' ').title())
            missing_fields.append(label)

    if not time_constraints: missing_fields.append(LABELS["Time Constraints"])

    if missing_fields:
        st.error("Please fill in all required fields:")
        for field in missing_fields:
            st.write(f"- {field}")
    else:
        # Prepare data for files
        # Serialize manuscripts for TSV
        ms_summary = []
        for ms in st.session_state.manuscript_list:
            ms_entry = {
                'category': ms['category'],
                'type': ms['type'],
                'value': ms['url'] if ms['type'] == 'URL' else ms['file'].name
            }
            ms_summary.append(ms_entry)

        all_responses = {
            "Proposal Type": proposal_type,
            "Scientific POC Name": sci_poc_name,
            "Scientific POC Email": sci_poc_email,
            "Scientific POC Phone": sci_poc_phone,
            "Technical POC Name": tech_poc_name,
            "Technical POC Email": tech_poc_email,
            "Technical POC Phone": tech_poc_phone,
            "Legal POC Name": legal_poc_name,
            "Legal POC Email": legal_poc_email,
            "Legal POC Phone": legal_poc_phone,
            "Time Constraints": time_constraints,
            "Title": title,
            "Nickname": nickname,
            "Authors": authors_raw,
            "Abstract": abstract,
            "Published Elsewhere": published_elsewhere,
            "Manuscripts": json.dumps(ms_summary)
        }
        all_responses.update(extra_data)

        # Generate TSV
        tsv_buffer = io.StringIO()
        # Drop complex objects if any remain, though here it should be fine
        df = pd.DataFrame([all_responses])
        df.to_csv(tsv_buffer, sep='\t', index=False)

        # Prepare Data for Reports (DOCX and PDF)
        report_data = [] # List of (label, value) pairs

        for key, value in all_responses.items():
            if key == "Manuscripts":
                label = "Manuscripts/Preprints"
                ms_text = ""
                for ms in st.session_state.manuscript_list:
                    val = ms['url'] if ms['type'] == 'URL' else ms['file'].name
                    ms_text += f"- [{ms['category']}] {val}\n"
                report_data.append((label, ms_text.strip()))
            elif "POC Phone" in key:
                if value and str(value).lower() != "nan" and str(value).strip():
                    label = LABELS.get(key, key)
                    report_data.append((label, str(value)))
            elif key == "data_components":
                report_data.append((LABELS[key], component_text(components)))
            elif key == "external_resources":
                report_data.append((LABELS[key], resource_text(resources) or "None provided"))
            elif key in {'label_catalog_version', 'proposal_model_version'}:
                continue  # Troubleshooting provenance belongs in the document footer.
            else:
                label = LABELS.get(key, key)
                if isinstance(value, list):
                    display_val = ", ".join(map(str, value))
                else:
                    display_val = str(value)
                report_data.append((label, display_val))

        # Generate DOCX
        doc = Document()
        from docx.shared import RGBColor
        from docx.oxml.ns import qn
        title_style = doc.styles['Title']
        title_style.font.color.rgb = RGBColor(0, 0, 0)
        title_style.font.underline = False
        for border in title_style.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
        doc.add_heading(f"TCIA Dataset Proposal: {title}", 0)

        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'

        for label, value in report_data:
            row_cells = table.add_row().cells
            row_cells[0].text = label
            row_cells[1].text = value
            # Make the first column bold
            for paragraph in row_cells[0].paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        doc.add_page_break()
        add_curation_placeholders(doc)
        from tcia_proposals.summary_document import add_reference_footer
        add_reference_footer(doc, all_responses)
        docx_buffer = io.BytesIO()
        doc.save(docx_buffer)
        docx_buffer.seek(0)

        # Generate PDF
        pdf_buffer = io.BytesIO()
        try:
            # Prepare Exhibit A pages using Platypus
            exhibit_a_buffer = io.BytesIO()

            # Custom canvas to add headers/footers
            class ExhibitCanvas(canvas.Canvas):
                def __init__(self, *args, **kwargs):
                    canvas.Canvas.__init__(self, *args, **kwargs)
                    self.pages = []

                def showPage(self):
                    self.pages.append(dict(self.__dict__))
                    self._startPage()

                def save(self):
                    page_count = len(self.pages)
                    for count, page in enumerate(self.pages):
                        self.__dict__.update(page)
                        self.draw_canvas(count+1, page_count)
                        canvas.Canvas.showPage(self)
                    canvas.Canvas.save(self)

                def draw_canvas(self, page_num, total_pages):
                    # Sequential numbering starts from 6
                    display_page = page_num + 5
                    # Total pages for the whole document (template has 7 pages, we replaced 1 with N)
                    total_display = total_pages + 6
                    draw_dsa_footer(self, display_page, total_display)

            doc_exhibit = SimpleDocTemplate(
                exhibit_a_buffer,
                pagesize=letter,
                rightMargin=50, leftMargin=50,
                topMargin=70, bottomMargin=50
            )

            styles = getSampleStyleSheet()
            style_center = ParagraphStyle(
                name='Center',
                parent=styles['Normal'],
                alignment=TA_CENTER
            )
            style_cell = ParagraphStyle(
                'TableCell',
                parent=styles['Normal'],
                fontSize=10,
                leading=12,
                wordWrap='LTR'
            )
            style_cell_bold = ParagraphStyle(
                'TableCellBold',
                parent=style_cell,
                fontName='Helvetica-Bold'
            )

            elements = []

            # Exhibit A Header
            elements.append(Paragraph("<br/><br/><br/><b><font size=14>EXHIBIT A</font></b>", style_center))
            elements.append(Paragraph("<b><font size=14>DESCRIPTION OF SUBMISSION DATA</font></b>", style_center))
            elements.append(Spacer(1, 20))

            # Table Data
            table_data = []
            for label, val in report_data:
                # Escape user input for ReportLab XML-like markup
                escaped_label = escape(label)
                escaped_val = escape(val).replace('\n', '<br/>')
                table_data.append([
                    Paragraph(f"<b>{escaped_label}</b>", style_cell),
                    Paragraph(escaped_val, style_cell)
                ])

            t = Table(table_data, colWidths=[150, 350])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            elements.append(t)

            doc_exhibit.build(elements, canvasmaker=ExhibitCanvas)
            exhibit_a_buffer.seek(0)

            # Merge with template
            reader = PyPDF2.PdfReader(AGREEMENT_TEMPLATE)
            writer = PyPDF2.PdfWriter()

            new_exhibit_reader = PyPDF2.PdfReader(exhibit_a_buffer)
            total_exhibit_pages = len(new_exhibit_reader.pages)
            total_final_pages = total_exhibit_pages + 6

            # Pages 1-5
            for i in range(5):
                page = replace_dsa_page_number(
                    reader.pages[i],
                    page_number=i + 1,
                    total_pages=total_final_pages,
                )
                writer.add_page(page)

            # Insert Exhibit A pages
            for page in new_exhibit_reader.pages:
                writer.add_page(page)

            # The template's final page is Exhibit B. Renumber it after the
            # variable-length Exhibit A pages have been inserted.
            last_page = replace_dsa_page_number(
                reader.pages[6],
                page_number=total_final_pages,
                total_pages=total_final_pages,
            )
            writer.add_page(last_page)

            writer.write(pdf_buffer)
            pdf_buffer.seek(0)
            pdf_success = True
        except Exception as e:
            st.error(f"Error generating PDF: {e}")
            pdf_success = False

        # Create ZIP
        today = datetime.date.today().isoformat()
        pkg_name = f"{nickname}_TCIA_Proposal_{today}.zip"
        summary_tsv_name = f"{nickname}_proposal_summary_{today}.tsv"
        docx_name = f"{nickname}_proposal_summary_{today}.docx"
        pdf_name = f"{nickname}_DSA_{today}.pdf"

        zip_buffer = io.BytesIO()
        with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
            zip_file.writestr(summary_tsv_name, tsv_buffer.getvalue())
            zip_file.writestr(docx_name, docx_buffer.getvalue())
            if pdf_success:
                zip_file.writestr(pdf_name, pdf_buffer.getvalue())
            for ms in st.session_state.manuscript_list:
                if ms['type'] == 'File':
                    zip_file.writestr(ms['file'].name, ms['file'].getvalue())
        zip_buffer.seek(0)

        # Store in session state for later use in email and to persist buttons
        st.session_state['proposal_files'] = {
            'tsv': tsv_buffer.getvalue(),
            'docx': docx_buffer.getvalue(),
            'pdf': pdf_buffer.getvalue() if pdf_success else None,
            'zip': zip_buffer.getvalue(),
            'title': title,
            'nickname': nickname,
            'proposal_type': proposal_type,
            'pocs': [sci_poc_email, tech_poc_email, legal_poc_email],
            'pkg_name': pkg_name
        }
        st.session_state['proposal_generated'] = True

        # Determine if we can send automatically
        can_send_auto = all([SMTP_SERVER, SMTP_USER, SMTP_PASSWORD])
        if can_send_auto:
            try:
                msg = MIMEMultipart()
                msg['From'] = SMTP_USER
                msg['To'] = HELP_DESK_EMAIL
                msg['Cc'] = ", ".join([sci_poc_email, tech_poc_email, legal_poc_email])
                msg['Subject'] = f"TCIA Dataset Proposal: {nickname}"

                body = f"Thank you for your interest in publishing your dataset on TCIA! Our Advisory Group meets quarterly to review proposals. We've received your proposal and will follow up shortly with any preliminary questions we have. If you have any questions in the meantime, please reply-all to this message.\n\n---\n\nType: {proposal_type}\nTitle: {title}\nSubmitters: {', '.join([sci_poc_email, tech_poc_email, legal_poc_email])}"
                msg.attach(MIMEText(body, 'plain'))

                part = MIMEBase('application', 'zip')
                part.set_payload(zip_buffer.getvalue())
                encoders.encode_base64(part)
                part.add_header('Content-Disposition', f'attachment; filename="{pkg_name}"')
                msg.attach(part)

                server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT)
                server.starttls()
                server.login(SMTP_USER, SMTP_PASSWORD)
                # Send to both To and Cc addresses
                recipients = [HELP_DESK_EMAIL, sci_poc_email, tech_poc_email, legal_poc_email]
                server.send_message(msg, to_addrs=recipients)
                server.quit()
                st.session_state['email_sent'] = True
            except Exception as e:
                st.error(f"Failed to send email: {e}")

if st.session_state.get('proposal_generated'):
    can_send_auto = all([SMTP_SERVER, SMTP_USER, SMTP_PASSWORD])
    files = st.session_state.get('proposal_files')

    if not can_send_auto:
        st.success("✅ Proposal documents generated successfully!")

        col1, col2, col3 = st.columns(3)
        today = datetime.date.today().isoformat()
        nickname_val = files.get('nickname', 'dataset')
        with col1:
            st.download_button("Download TSV (Import to Remapper)", data=files['tsv'], file_name=f"{nickname_val}_proposal_{today}.tsv", mime="text/tab-separated-values")
        with col2:
            st.download_button("Download DOCX Summary", data=files['docx'], file_name=f"{nickname_val}_proposal_summary_{today}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        with col3:
            if files['pdf']:
                st.download_button("Download PDF Agreement", data=files['pdf'], file_name=f"{nickname_val}_DSA_{today}.pdf", mime="application/pdf")

        st.info("💡 **Please download these documents and keep a copy for your records.**")
        st.info(f"💡 **Manual Submission**: Please email the generated files as attachments to **{HELP_DESK_EMAIL}**.")
    else:
        if st.session_state.get('email_sent'):
            st.success(f"📨 Proposal sent to {HELP_DESK_EMAIL} and CC'd to POCs!")
        else:
            st.info("Proposal generated but failed to send email. Please check your SMTP settings or contact support.")

# Footer
st.markdown("---")
st.markdown(f"""
<div style='text-align: center; color: gray; font-size: 0.9em;'>
TCIA Dataset Proposal Form | {HELP_DESK_EMAIL}
</div>
""", unsafe_allow_html=True)
