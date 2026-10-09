"""Import and export CRDC Data Hub submission-request workbooks."""

from __future__ import annotations

import ast
import datetime as dt
import re
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


CRDC_TEMPLATE_PATH = (
    Path(__file__).resolve().parent
    / "resources"
    / "crdc_submission_request_template_v3.5.xlsx"
)
CRDC_TEMPLATE_SHA256 = "0dc32ce2a6f73c0d55bd4ff594014d68143c080f0f328dc57c28cf510f62abe8"

REQUIRED_HEADERS = {
    "PI and Contact": "Principal Investigator First name",
    "Program and Study": "Study Title",
    "Data Access and Disease": "Number of subjects included in the submission",
    "Data Types": "Targeted Data Submission Delivery Date",
}


class CRDCWorkbookError(ValueError):
    """Raised when an uploaded workbook is not a supported CRDC proposal."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _unique(values: Iterable[Any]) -> list[str]:
    result = []
    seen = set()
    for value in values:
        cleaned = _text(value)
        if cleaned and cleaned not in seen:
            result.append(cleaned)
            seen.add(cleaned)
    return result


def _pipe_values(values: Iterable[Any]) -> list[str]:
    expanded = []
    for value in values:
        expanded.extend(part.strip() for part in _text(value).split("|") if part.strip())
    return _unique(expanded)


def _load(source: Any, *, data_only: bool = False):
    if isinstance(source, (bytes, bytearray)):
        source = BytesIO(source)
    elif hasattr(source, "seek"):
        source.seek(0)
    return load_workbook(source, data_only=data_only)


def _headers(ws) -> dict[str, int]:
    return {_text(cell.value): cell.column for cell in ws[1] if _text(cell.value)}


def _values(ws, header: str, *, max_row: int = 100) -> list[Any]:
    column = _headers(ws).get(header)
    if not column:
        return []
    return [ws.cell(row, column).value for row in range(2, max_row + 1)]


def _first(ws, header: str) -> Any:
    values = _values(ws, header, max_row=2)
    return values[0] if values else None


def _set(ws, header: str, value: Any, row: int = 2) -> None:
    column = _headers(ws).get(header)
    if not column:
        raise CRDCWorkbookError(f"CRDC template is missing column: {header}")
    ws.cell(row, column).value = value if value not in ("", [], ()) else None


def _split_name(name: Any) -> tuple[str, str]:
    cleaned = _text(name)
    if not cleaned:
        return "", ""
    if "," in cleaned:
        last, first = (part.strip() for part in cleaned.split(",", 1))
        return first, last
    parts = cleaned.split()
    if len(parts) == 1:
        return parts[0], ""
    return " ".join(parts[:-1]), parts[-1]


def _slug(value: Any) -> str:
    slug = re.sub(r"[^A-Za-z0-9-]+", "-", _text(value)).strip("-")
    return slug[:30] or "CRDC-proposal"


def _parse_int(value: Any) -> int | None:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value)
    match = re.search(r"\d[\d,]*", _text(value))
    return int(match.group().replace(",", "")) if match else None


def _yes(value: Any) -> bool:
    return _text(value).casefold() == "yes"


def _author_text(investigators: list[dict[str, Any]]) -> str:
    authors = []
    for investigator in investigators:
        name = ", ".join(
            part
            for part in (
                _text(investigator.get("last_name")),
                _text(investigator.get("first_name")),
            )
            if part
        )
        orcid = _text(investigator.get("person_orcid"))
        if name:
            authors.append(f"{name} - {orcid}" if orcid else name)
    return "; ".join(authors)


def is_crdc_workbook(source: Any) -> bool:
    """Return True when a workbook has the supported CRDC sheet contract."""
    try:
        workbook = _load(source, data_only=False)
    except Exception:
        return False
    try:
        for sheet_name, required_header in REQUIRED_HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                return False
            if required_header not in _headers(workbook[sheet_name]):
                return False
        return True
    finally:
        workbook.close()


def import_crdc_workbook(
    source: Any, *, proposal_type: str = "New Collection Proposal"
) -> dict[str, Any]:
    """Convert a CRDC submission workbook into remapper session payloads."""
    workbook = _load(source, data_only=True)
    try:
        for sheet_name, required_header in REQUIRED_HEADERS.items():
            if sheet_name not in workbook.sheetnames:
                raise CRDCWorkbookError(f"Missing CRDC worksheet: {sheet_name}")
            if required_header not in _headers(workbook[sheet_name]):
                raise CRDCWorkbookError(
                    f"Worksheet '{sheet_name}' is missing column: {required_header}"
                )

        contacts = workbook["PI and Contact"]
        study = workbook["Program and Study"]
        access = workbook["Data Access and Disease"]
        data_types = workbook["Data Types"]

        pi_first = _text(_first(contacts, "Principal Investigator First name"))
        pi_last = _text(_first(contacts, "Principal Investigator Last name"))
        investigators = []
        if pi_first or pi_last:
            investigators.append(
                {
                    "author_order": 1,
                    "first_name": pi_first,
                    "last_name": pi_last,
                    "middle_name": "",
                    "email": _text(_first(contacts, "Principal Investigator Email")),
                    "organization_name": _text(
                        _first(contacts, "Principal Investigator Institution")
                    ),
                    "person_orcid": _text(
                        _first(contacts, "Principal Investigator ORCID")
                    ),
                    "role": _text(_first(contacts, "Principal Investigator Position")),
                    "ror_id": "",
                }
            )

        same_contact = _yes(
            _first(contacts, "Primary Contact Same as Principal Investigator")
        )
        if same_contact:
            scientific_name = " ".join(part for part in (pi_first, pi_last) if part)
            scientific_email = _text(_first(contacts, "Principal Investigator Email"))
            scientific_phone = ""
        else:
            scientific_name = " ".join(
                part
                for part in (
                    _text(_first(contacts, "Primary Contact First name")),
                    _text(_first(contacts, "Primary Contact Last name")),
                )
                if part
            )
            scientific_email = _text(_first(contacts, "Primary Contact Email"))
            scientific_phone = _text(_first(contacts, "Primary Contact Phone number"))

        title = _text(_first(study, "Study Title"))
        nickname = (
            _text(_first(study, "Study Abbreviation"))
            or _text(_first(study, "Study ID"))
            or _slug(title)
        )
        description = _text(_first(study, "Study Description"))
        cancer_types = _pipe_values(_values(access, "Cancer Types"))
        cancer_types += [
            value
            for value in _pipe_values(_values(access, "Other cancer type(s)"))
            if value not in cancer_types
        ]

        selected_data_types = []
        type_map = {
            "Clinical": "Clinical",
            "Genomics": "Genomics",
            "Proteomics": "Proteomics",
        }
        for header, tcia_value in type_map.items():
            if _yes(_first(data_types, header)):
                selected_data_types.append(tcia_value)
        other_data_types = _pipe_values(_values(data_types, "Other Data Type(s)"))
        selected_data_types.extend(
            value for value in other_data_types if value not in selected_data_types
        )

        file_types = _values(data_types, "File Type")
        file_extensions = _values(data_types, "File Extension")
        file_formats = []
        for file_type, extension in zip(file_types, file_extensions):
            if _text(file_type) or _text(extension):
                file_formats.append(
                    " - ".join(part for part in (_text(file_type), _text(extension)) if part)
                )

        funding_agencies = _pipe_values(
            _values(study, "Funding Agency/Organization")
        )
        grants = _pipe_values(_values(study, "Grant or Contract Number(s)"))

        proposal_data = {
            "Proposal Type": proposal_type,
            "Scientific POC Name": scientific_name,
            "Scientific POC Email": scientific_email,
            "Scientific POC Phone": scientific_phone,
            "Title": title,
            "Nickname": nickname,
            "Authors": _author_text(investigators),
            "Abstract": description,
            "disease_site": [],
            "diagnosis": cancer_types,
            "image_types": ["Other"] if _yes(_first(data_types, "Imaging")) else [],
            "supporting_data": selected_data_types,
            "file_formats": "; ".join(file_formats),
            "disk_space": " | ".join(
                _unique(_values(data_types, "Estimated data size"))
            ),
            "number_of_subjects": _parse_int(
                _first(access, "Number of subjects included in the submission")
            )
            or "",
            "funding_agency": " | ".join(funding_agencies),
            "funding_source_program_name": _text(_first(study, "Program")),
            "grant_id": " | ".join(grants),
            "Published Elsewhere": "",
            "Time Constraints": "",
            "acknowledgements": "",
            "crdc_source": True,
        }

        program_value = _text(_first(study, "Program"))
        if program_value == "Other":
            program_name = _text(_first(study, "Program Title")) or "Other"
            program_short_name = (
                _text(_first(study, "Program Abbreviation")) or _slug(program_name)
            )
            program_description = _text(_first(study, "Program Description"))
        else:
            program_name = program_value or "Community"
            program_short_name = program_value or "Community"
            program_description = ""
            if program_value and "ProgramList" in workbook.sheetnames:
                for row in workbook["ProgramList"].iter_rows(values_only=True):
                    display = _text(row[4] if len(row) > 4 else "")
                    if display == program_value:
                        program_name = _text(row[1]) or display
                        program_short_name = _text(row[2]) or display
                        program_description = _text(row[3])
                        break

        program = {
            "program_name": program_name,
            "program_short_name": program_short_name,
            "institution_name": "",
            "program_short_description": program_description[:255],
            "program_full_description": program_description,
            "program_external_url": "",
        }

        related_works = []
        headers = _headers(study)
        for row in range(2, 101):
            doi = _text(study.cell(row, headers.get("DOI", 0)).value) if headers.get("DOI") else ""
            pmid = (
                _text(study.cell(row, headers.get("PubMed ID (PMID)", 0)).value)
                if headers.get("PubMed ID (PMID)")
                else ""
            )
            for header, relationship in (
                ("Existing Publication Title", "IsDerivedFrom"),
                ("Planned Publication Title", "Describes"),
            ):
                column = headers.get(header)
                work_title = _text(study.cell(row, column).value) if column else ""
                if work_title:
                    related_works.append(
                        {
                            "title": work_title,
                            "publication_type": "Journal Article",
                            "relationship_type": relationship,
                            "authorship": "",
                            "DOI": doi if relationship == "IsDerivedFrom" else "",
                            "resource_identifier": (
                                f"PMID:{pmid}"
                                if pmid and relationship == "IsDerivedFrom"
                                else ""
                            ),
                        }
                    )

        return {
            "proposal_data": proposal_data,
            "program": program,
            "investigators": investigators,
            "related_works": related_works,
            "cicadas": {"abstract": description},
        }
    finally:
        workbook.close()


def _metadata_first(metadata: dict[str, Any], entity: str) -> dict[str, Any]:
    values = metadata.get(entity) or []
    return values[0] if values else {}


def _join_for_cell(value: Any) -> str:
    if isinstance(value, (list, tuple, set)):
        return " | ".join(_unique(value))
    return _text(value)


def _list_value(value: Any) -> list[str]:
    if isinstance(value, (list, tuple, set)):
        return _unique(value)
    text = _text(value)
    if not text:
        return []
    if text.startswith("[") and text.endswith("]"):
        try:
            parsed = ast.literal_eval(text)
            if isinstance(parsed, (list, tuple, set)):
                return _unique(parsed)
        except (SyntaxError, ValueError):
            pass
    return _pipe_values([text])


def _known_values(workbook, sheet_name: str, column: int) -> set[str]:
    cells = next(
        workbook[sheet_name].iter_cols(
            min_col=column, max_col=column, values_only=False
        )
    )
    return {
        _text(cell.value)
        for cell in cells
        if _text(cell.value)
    }


def build_crdc_workbook(
    proposal_data: dict[str, Any],
    metadata: dict[str, Any],
    cicadas: dict[str, Any] | None = None,
    *,
    template_path: Path | str = CRDC_TEMPLATE_PATH,
) -> tuple[bytes, list[str]]:
    """Populate the canonical CRDC v3.5 template without changing its layout."""
    workbook = load_workbook(template_path)
    warnings = []
    cicadas = cicadas or {}

    dataset = _metadata_first(metadata, "Dataset")
    program = _metadata_first(metadata, "Program")
    investigators = sorted(
        metadata.get("Investigator") or [],
        key=lambda item: int(item.get("author_order", 999))
        if str(item.get("author_order", "")).isdigit()
        else 999,
    )
    related_works = metadata.get("Related_Work") or []

    metadata_ws = workbook["Metadata"]
    _set(metadata_ws, "Form Version", "3.5")
    _set(metadata_ws, "Export Date", dt.datetime.now(dt.timezone.utc).isoformat())

    contacts = workbook["PI and Contact"]
    pi = investigators[0] if investigators else {}
    _set(contacts, "Principal Investigator First name", pi.get("first_name"))
    _set(contacts, "Principal Investigator Last name", pi.get("last_name"))
    _set(contacts, "Principal Investigator Position", pi.get("role"))
    _set(contacts, "Principal Investigator Email", pi.get("email"))
    _set(contacts, "Principal Investigator ORCID", pi.get("person_orcid"))
    pi_institution = _text(pi.get("organization_name"))
    _set(contacts, "Principal Investigator Institution", pi_institution)
    known_institutions = _known_values(workbook, "InstitutionList", 2)
    if pi_institution and pi_institution not in known_institutions:
        warnings.append(
            "The principal investigator institution is not in the CRDC dropdown list; review it before submission."
        )

    scientific_name = proposal_data.get("Scientific POC Name")
    scientific_first, scientific_last = _split_name(scientific_name)
    same_pi = bool(pi) and (
        scientific_first.casefold() == _text(pi.get("first_name")).casefold()
        and scientific_last.casefold() == _text(pi.get("last_name")).casefold()
        and _text(proposal_data.get("Scientific POC Email")).casefold()
        == _text(pi.get("email")).casefold()
    )
    _set(contacts, "Primary Contact Same as Principal Investigator", "Yes" if same_pi else "No")
    if not same_pi:
        _set(contacts, "Primary Contact First name", scientific_first)
        _set(contacts, "Primary Contact Last name", scientific_last)
        _set(contacts, "Primary Contact Email", proposal_data.get("Scientific POC Email"))
        _set(contacts, "Primary Contact Phone number", proposal_data.get("Scientific POC Phone"))

    additional_contacts = []
    for prefix in ("Technical POC", "Legal POC"):
        name = _text(proposal_data.get(f"{prefix} Name"))
        email = _text(proposal_data.get(f"{prefix} Email"))
        phone = _text(proposal_data.get(f"{prefix} Phone"))
        if name or email or phone:
            first, last = _split_name(name)
            additional_contacts.append((first, last, prefix, email, phone))
    for row, (first, last, position, email, phone) in enumerate(additional_contacts, 2):
        _set(contacts, "Additional Contact(s) First name", first, row)
        _set(contacts, "Additional Contact(s) Last name", last, row)
        _set(contacts, "Additional Contact(s) Position", position, row)
        _set(contacts, "Additional Contact(s) Email", email, row)
        _set(contacts, "Additional Contact(s) Phone number", phone, row)

    study = workbook["Program and Study"]
    program_candidates = {
        _text(program.get("program_name")),
        _text(program.get("program_short_name")),
    } - {""}
    program_display = ""
    for row in workbook["ProgramList"].iter_rows(values_only=True):
        if program_candidates.intersection(_text(value) for value in row[:5]):
            program_display = _text(row[4])
            break
    if program_display:
        _set(study, "Program", program_display)
    else:
        _set(study, "Program", "Other")
        _set(study, "Program Title", program.get("program_name"))
        _set(study, "Program Abbreviation", program.get("program_short_name"))
        _set(
            study,
            "Program Description",
            program.get("program_full_description")
            or program.get("program_short_description"),
        )

    _set(study, "Study Title", dataset.get("dataset_long_name") or proposal_data.get("Title"))
    _set(study, "Study Abbreviation", dataset.get("dataset_short_name") or proposal_data.get("Nickname"))
    _set(
        study,
        "Study Description",
        cicadas.get("abstract")
        or dataset.get("dataset_abstract")
        or proposal_data.get("Abstract"),
    )
    funding_agencies = _list_value(
        dataset.get("funding_agency") or proposal_data.get("funding_agency")
    )
    grant_ids = _list_value(dataset.get("grant_id") or proposal_data.get("grant_id"))
    known_funding_agencies = _known_values(workbook, "FundingAgencyList", 1)
    for row, agency in enumerate(funding_agencies, 2):
        _set(study, "Funding Agency/Organization", agency, row)
        if agency not in known_funding_agencies:
            warnings.append(
                f"Funding agency '{agency}' is not in the CRDC dropdown list; review it before submission."
            )
    for row, grant_id in enumerate(grant_ids, 2):
        _set(study, "Grant or Contract Number(s)", grant_id, row)
    _set(study, "Repository Name", "The Cancer Imaging Archive")
    _set(study, "Study ID", dataset.get("dataset_short_name") or proposal_data.get("Nickname"))

    existing_row = 2
    planned_row = 2
    for work in related_works:
        title = _text(work.get("title"))
        if not title:
            continue
        if work.get("relationship_type") == "Describes":
            _set(study, "Planned Publication Title", title, planned_row)
            planned_row += 1
        else:
            _set(study, "Existing Publication Title", title, existing_row)
            row = existing_row
            existing_row += 1
            _set(study, "DOI", work.get("DOI"), row)
            identifier = _text(work.get("resource_identifier"))
            if identifier.casefold().startswith("pmid:"):
                _set(study, "PubMed ID (PMID)", identifier.split(":", 1)[1], row)

    supporting = _list_value(proposal_data.get("supporting_data"))
    image_types = _list_value(proposal_data.get("image_types"))
    derived_types = _list_value(proposal_data.get("derived_types"))
    data_type_tokens = {value.casefold() for value in supporting}
    imaging = bool(image_types or derived_types) or proposal_data.get("Proposal Type") in {
        "New Collection Proposal",
        "Analysis Results Proposal",
    }
    submitted = ["imaging"] if imaging else []
    for key, label in (("clinical", "clinicalTrial"), ("genomics", "genomics"), ("proteomics", "proteomics")):
        if key in data_type_tokens:
            submitted.append(label)
    other_submitted = [
        value
        for value in supporting
        if value.casefold() not in {"clinical", "genomics", "proteomics", "no additional data"}
    ]
    _set(study, "Data Type(s) Submitted", " | ".join(submitted))
    _set(study, "Other Data Type(s)", " | ".join(other_submitted))

    access = workbook["Data Access and Disease"]
    open_access = proposal_data.get("Access Types: Open Access")
    controlled_access = proposal_data.get("Access Types: Controlled Access")
    _set(access, "Access Types: Open Access", open_access)
    _set(access, "Access Types: Controlled Access", controlled_access)
    diagnoses = _list_value(proposal_data.get("diagnosis"))
    known_cancers = _known_values(workbook, "CancerTypeList", 1)
    recognized = [value for value in diagnoses if value in known_cancers]
    unrecognized = [value for value in diagnoses if value not in known_cancers]
    for row, value in enumerate(recognized, 2):
        _set(access, "Cancer Types", value, row)
    _set(access, "Other cancer type(s)", " | ".join(unrecognized))
    subject_count = _parse_int(proposal_data.get("number_of_subjects"))
    _set(access, "Number of subjects included in the submission", subject_count)

    data_ws = workbook["Data Types"]
    _set(data_ws, "Clinical", "Yes" if "clinical" in data_type_tokens else "No")
    _set(data_ws, "Genomics", "Yes" if "genomics" in data_type_tokens else "No")
    _set(data_ws, "Imaging", "Yes" if imaging else "No")
    _set(data_ws, "Proteomics", "Yes" if "proteomics" in data_type_tokens else "No")
    _set(data_ws, "Other Data Type(s)", " | ".join(other_submitted))

    clinical_subtypes = {
        "Demographic Data": "demographic",
        "Relapse/Recurrence Data": "relapse",
        "Diagnosis Data": "diagnosis",
        "Outcome Data": "outcome",
        "Treatment Data": "treatment",
        "Biospecimen Data": "biospecimen",
    }
    supporting_folded = " ".join(supporting).casefold()
    for header, token in clinical_subtypes.items():
        if token in supporting_folded:
            _set(data_ws, header, "Yes")

    format_text = _text(proposal_data.get("file_formats"))
    format_entries = [entry.strip() for entry in format_text.split(";") if entry.strip()]
    known_extensions = _known_values(workbook, "FileTypeList", 2)
    for row, entry in enumerate(format_entries, 2):
        source_type, separator, extension = entry.partition(" - ")
        lowered = source_type.casefold()
        if "clinical" in lowered:
            crdc_file_type = "Clinical data"
        elif "protein" in lowered or "proteom" in lowered:
            crdc_file_type = "Protein expression data"
        elif "sequence" in lowered or "genom" in lowered:
            crdc_file_type = "Derived sequencing data"
        else:
            crdc_file_type = "Imaging data"
        exported_extension = extension if separator else source_type
        _set(data_ws, "File Type", crdc_file_type, row)
        _set(data_ws, "File Extension", exported_extension, row)
        if exported_extension and exported_extension not in known_extensions:
            warnings.append(
                f"File extension '{exported_extension}' is not in the CRDC dropdown list; review it before submission."
            )
        if row == 2:
            _set(data_ws, "Estimated data size", proposal_data.get("disk_space"), row)

    comment_parts = []
    for label, key in (
        ("Time constraints", "Time Constraints"),
        ("Published elsewhere", "Published Elsewhere"),
        ("TCIA rationale", "why_tcia"),
    ):
        value = _join_for_cell(proposal_data.get(key))
        if value:
            comment_parts.append(f"{label}: {value}")
    comments = "\n".join(comment_parts)
    if len(comments) > 500:
        warnings.append(
            "CRDC additional comments were limited to 500 characters; review that field before submission."
        )
        comments = comments[:500]
    _set(data_ws, "Additional Comments or Information about this submission.", comments)

    if not open_access and not controlled_access:
        warnings.append(
            "CRDC access-type fields were left blank because the TCIA proposal does not establish open versus controlled access."
        )
    warnings.append(
        "Review CRDC-only required fields before submission, including dates, dbGaP status, file counts, and de-identification confirmations."
    )

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue(), warnings
