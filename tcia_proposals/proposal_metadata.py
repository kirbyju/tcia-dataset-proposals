"""Dataset proposal metadata shared by the form, exports, and remapper."""

import ast
import json
from pathlib import Path
from urllib.parse import urlsplit
from urllib.request import urlopen

CATALOG_PATH = Path(__file__).parent / "resources/label_catalog.json"
PROPOSAL_MODEL_PATH = Path(__file__).parent / "resources/proposal_data_model.json"
DATA_MODELS_URL = "https://kirbyju.github.io/tcia-dataset-proposals/labels/"
_last_good_reference = None
FUNDING_PROPERTIES = (
    "funderName", "funderIdentifier", "funderIdentifierType",
    "awardNumber", "awardURI", "awardTitle",
)


def load_catalog():
    return json.loads(CATALOG_PATH.read_text(encoding="utf-8"))


def load_proposal_model():
    return json.loads(PROPOSAL_MODEL_PATH.read_text(encoding="utf-8"))


def load_form_reference():
    """Use the published atomic catalog/rules pair, retaining a bundled fallback."""
    global _last_good_reference
    bundled = {'catalog': load_catalog(), 'proposal_model': load_proposal_model()}
    def version(value):
        if not isinstance(value, str):
            raise ValueError('Invalid guidance version')
        return tuple(int(part) for part in value.split('.'))
    fallback = _last_good_reference or bundled
    if version(fallback['proposal_model']['version']) < version(bundled['proposal_model']['version']):
        fallback = bundled
    try:
        with urlopen(DATA_MODELS_URL + 'form-reference.json', timeout=3) as response:
            live = json.loads(response.read(8_000_001))
        catalog, model = live['catalog'], live['proposal_model']
        expected = {'disease_type', 'disease_location', 'data_category', 'data_type', 'file_type', 'species', 'program', 'external_resource'}
        if set(catalog['fields']) != expected or catalog.get('wordpress_registered_terms_complete') is not True:
            raise ValueError('Incomplete public registry capture')
        if catalog.get('approved_vocabulary') is not False or not catalog.get('sources'):
            raise ValueError('Missing public provenance')
        for field, rows in catalog['fields'].items():
            labels = [row['label'] for row in rows]
            if not rows or len(labels) != len(set(labels)) or any(not isinstance(label, str) or not label.strip() for label in labels):
                raise ValueError(f'Invalid labels in {field}')
        if version(model['version']) < version(fallback['proposal_model']['version']):
            raise ValueError('Published selection guidance is older than the application')
        if catalog['generated_at_utc'] < fallback['catalog']['generated_at_utc']:
            raise ValueError('Published capture is older than the bundled capture')
        if not all(isinstance(model[key], dict) for key in ('category_data_types', 'category_aliases', 'data_type_explanations')):
            raise ValueError('Invalid selection guidance')
        if model['category_data_types'].get('Other') != ['Other']:
            raise ValueError('Other category must offer only Other type')
        if not all(isinstance(model[key], list) for key in ('category_order', 'species_order')):
            raise ValueError('Invalid display order')
        _last_good_reference = live
        return live
    except (OSError, ValueError, KeyError, TypeError):
        return fallback


def species_options(catalog, model=None):
    labels = label_options(catalog, 'species')
    order = (model or load_proposal_model())['species_order']
    return [value for value in order if value in labels] + [value for value in labels if value not in order]


def category_options(catalog, model=None):
    registered = set(label_options(catalog, 'data_category'))
    return [value for value in (model or load_proposal_model())['category_order'] if value in registered]


def data_type_options(catalog, category, model=None):
    model = model or load_proposal_model()
    category = model['category_aliases'].get(category, category)
    allowed = model['category_data_types'].get(category, [])
    return [value for value in label_options(catalog, 'data_type') if '*' in allowed or value in allowed]


def data_type_display(value, model=None):
    explanation = (model or load_proposal_model())['data_type_explanations'].get(value)
    return f'{value} ({explanation})' if explanation else value


def normalize_resource_url(value):
    """Add HTTPS to a bare host, preserving explicit schemes for validation."""
    value = value.strip()
    if value.startswith('//'):
        return 'https:' + value
    if not value or '://' in value or value.lower().startswith(('mailto:', 'javascript:', 'data:', 'file:')):
        return value
    try:
        parsed = urlsplit('https://' + value)
    except ValueError:
        return value
    if parsed.hostname and '.' in parsed.hostname and not any(c.isspace() for c in value) and not parsed.username:
        return 'https://' + value
    return value


def label_options(catalog, field):
    """Only currently registered choices; retain exact spelling and scope."""
    return sorted({row["label"] for row in catalog["fields"][field]
                   if row.get("registry_membership") == "registered"
                   and row.get("confirmed_decision", {}).get("decision") != "remove from field"},
                  key=str.casefold)


def parse_rows(value):
    if isinstance(value, list):
        return value
    if not isinstance(value, str) or not value.strip():
        return []
    try:
        result = json.loads(value)
    except (ValueError, TypeError):
        try:
            result = ast.literal_eval(value)  # Older TSV packages used Python list strings.
        except (ValueError, SyntaxError):
            return []
    return result if isinstance(result, list) else []


def validate_components(rows, model=None):
    errors = []
    if not rows:
        return ["Add at least one data component."]
    for index, row in enumerate(rows, 1):
        if not isinstance(row.get("category"), str) or not row["category"].strip():
            errors.append(f"Data component {index}: choose one category.")
        types = row.get("data_types")
        if not isinstance(types, list) or not types or any(not isinstance(t, str) or not t.strip() for t in types):
            errors.append(f"Data component {index}: choose at least one data type.")
        if not isinstance(row.get("file_format"), str) or not row["file_format"].strip():
            errors.append(f"Data component {index}: choose one file format.")
        for key in ("category", "data_types", "file_format"):
            selected = row.get(key)
            has_other = "Other" in selected if isinstance(selected, list) else selected == "Other"
            if has_other and not row.get(f"other_{key}", "").strip():
                errors.append(f"Data component {index}: describe the other {key.replace('_', ' ')}.")
        category = row.get('category')
        if isinstance(category, str) and isinstance(types, list):
            selection_model = model or load_proposal_model()
            allowed = selection_model['category_data_types'].get(selection_model['category_aliases'].get(category, category), [])
            if allowed and '*' not in allowed and any(value not in allowed for value in types):
                errors.append(f"Data component {index}: choose data types that fit {category}, or describe Other.")
    return errors


def validate_resources(rows):
    errors = []
    for index, row in enumerate(rows, 1):
        if not row.get("category"):
            errors.append(f"External resource {index}: choose a category.")
        if row.get("category") == "Other" and not row.get("other_category", "").strip():
            errors.append(f"External resource {index}: describe the other category.")
        try:
            url = urlsplit(row.get("url", ""))
            valid = url.scheme in {"http", "https"} and bool(url.hostname) and not any(c.isspace() for c in row.get("url", ""))
        except ValueError:
            valid = False
        if not valid:
            errors.append(f"External resource {index}: enter a complete http or https URL.")
    return errors


def component_text(rows):
    return "\n".join(
        f"{i}. {row['category']} — {', '.join(row['data_types'])} — {row['file_format']}"
        + "".join(f"; {key.removeprefix('other_').replace('_', ' ')}: {value}"
                  for key, value in row.items() if key.startswith("other_") and value)
        for i, row in enumerate(rows, 1)
    )


def resource_text(rows):
    return "\n".join(f"{row['category']}" + (f" ({row['other_category']})" if row.get("other_category") else "")
                     + f": {row['url']}" + (f" -- {row['context']}" if row.get("context") else "") for row in rows)


def legacy_component_fields(rows, analysis=False):
    """Compatibility summaries; structured rows preserve every type/format association."""
    images, supporting, formats = [], [], []
    for row in rows:
        image = row["category"] in {"Radiology Images", "Pathology Images", "Histopathology Images"}
        target = images if image else supporting
        if row["category"] == "Clinical Data":
            supporting.append("Clinical")
        target.extend(row["data_types"])
        formats.extend(f"{row['category']}: {data_type} - {row['file_format']}" for data_type in row["data_types"])
    return {
        "image_types": list(dict.fromkeys(images)),
        "supporting_data": list(dict.fromkeys(supporting)),
        "derived_types": list(dict.fromkeys(t for r in rows for t in r["data_types"])) if analysis else [],
        "file_formats": "; ".join(formats),
    }


def add_curation_placeholders(doc):
    """Blank fields for later curation, in the submitter's funding table layout."""
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    doc.add_heading("Funding information", level=1)
    doc.add_paragraph("To be completed during curation. Add columns for additional funding sources as needed.")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    section = doc.sections[-1]
    available = (section.page_width - section.left_margin - section.right_margin) / Inches(1)
    widths = (available * 0.34, available * 0.33, available * 0.33)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, text in zip(table.rows[0].cells, ("DataCite Property", "Funding Source 1", "Funding Source 2")):
        cell.text = text
        for run in cell.paragraphs[0].runs:
            run.bold = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for field in FUNDING_PROPERTIES:
        cells = table.add_row().cells
        cells[0].text = field
        cells[0].paragraphs[0].runs[0].bold = True
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    doc.add_heading("Acknowledgements", level=1)
    doc.add_paragraph("To be completed during curation.")
    blank = doc.add_table(rows=1, cols=1)
    blank.style = "Table Grid"
    blank.rows[0].height = Inches(0.65)
