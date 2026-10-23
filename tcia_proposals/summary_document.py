"""Word summary formatting and exact proposal-label checks."""

import ast
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from tcia_proposals.proposal_metadata import FUNDING_PROPERTIES, label_options, load_catalog, parse_rows

def review_term(value, proposed=''):
    """Mark explicit Other selections for FNL without changing stored terms."""
    if value == 'Other':
        return f"Other -- Requires FNL Review -- {proposed or '<proposed term>'}"
    if isinstance(value, str) and value.startswith('Other - '):
        return f"Other -- Requires FNL Review -- {value[len('Other - '):]}"
    return value


def tracking_component_text(rows):
    blocks = []
    for row in rows:
        category = review_term(row['category'], row.get('other_category'))
        types = [review_term(value, row.get('other_data_types')) for value in row['data_types']]
        file_format = review_term(row['file_format'], row.get('other_file_format'))
        type_label = 'Type' if row['data_types'] == ['Other'] else 'Type(s)'
        blocks.append(f"Category: {category}\n{type_label}: {', '.join(types)}\nFormat: {file_format}")
    return '\n\n'.join(blocks)


def add_reference_footer(doc, proposal_data):
    from docx.shared import RGBColor
    values = [("Data model reference version", proposal_data.get('label_catalog_version')),
              ("Proposal selection guidance version", proposal_data.get('proposal_model_version'))]
    values = [(label, value) for label, value in values if value]
    if not values:
        return
    paragraph = doc.sections[-1].footer.paragraphs[0]
    paragraph.clear()
    for index, (label, value) in enumerate(values):
        if index:
            paragraph.add_run('\n')
        paragraph.add_run(f'{label}: ').bold = True
        paragraph.add_run(str(value))
    for run in paragraph.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string('666666')


def add_link(paragraph, text, url):
    """Write a clickable Word hyperlink rather than literal Markdown."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2467A8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    label = OxmlElement("w:t")
    label.text = text
    run.append(label)
    link.append(run)
    paragraph._p.append(link)


def rights_placeholder(cell):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    add_link(paragraph, "CC BY 4.0", "https://creativecommons.org/licenses/by/4.0/")
    paragraph.add_run(" or ")
    add_link(paragraph, "NIH Controlled Data Access Policy",
             "https://www.cancerimagingarchive.net/nih-controlled-data-access-policy/")


def exact_dataset_labels(value, field, catalog=None):
    """Preserve exact labels or explicit disease/site Other answers; never fuzzy-map."""
    if value is None or (isinstance(value, float) and value != value):
        return []
    if isinstance(value, str) and not value.strip():
        return []
    values = parse_rows(value)
    if not values and (isinstance(value, list) or (isinstance(value, str) and value.replace(" ", "") == "[]")):
        return []
    if not values:
        values = [value]
    allowed = set(label_options(catalog or load_catalog(), field))
    for label in values:
        custom = (field in {"disease_type", "disease_location"} and isinstance(label, str)
                  and label.startswith("Other - ") and bool(label[len("Other - "):].strip()))
        if not isinstance(label, str) or (label not in allowed and not custom):
            raise ValueError(
                f"Unrecognized {field.replace('_', ' ')} label: {label!r}. "
                "Use an exact TCIA Data Models label"
                + (" or an explicit 'Other - description' answer." if field != "species" else ".")
            )
    return values


def add_funding_row(table, doc, sources=None):
    """A full-width funding table within its own Dataset Information row."""
    sources = sources or [{}, {}]
    sources = list(sources) + [{}] * max(0, 2 - len(sources))
    funding_row = table.add_row()
    funding_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cells = funding_row.cells
    cell = cells[0].merge(cells[1])
    cell.text = "Funding"
    cell.paragraphs[0].runs[0].bold = True
    cell.add_paragraph()  # Blank line before the funding example.
    funding_cell = cell
    nested = cell.add_table(rows=1, cols=len(sources) + 1)
    nested.style = "Table Grid"
    nested.autofit = False
    section = doc.sections[-1]
    width = (section.page_width - section.left_margin - section.right_margin) / Inches(1) - 0.16
    widths = [width * 0.34] + [width * 0.66 / len(sources)] * len(sources)
    for column, col_width in zip(nested.columns, widths):
        column.width = Inches(col_width)
    for cell, title in zip(nested.rows[0].cells,
                           ["Datacite Property"] + [f"Funding Source {i}" for i in range(1, len(sources) + 1)]):
        cell.text = title
        cell.paragraphs[0].runs[0].bold = True
    nested.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for field in FUNDING_PROPERTIES:
        row = nested.add_row()
        row.cells[0].text = field
        row.cells[0].paragraphs[0].runs[0].bold = True
        for target, source in zip(row.cells[1:], sources):
            target.text = str(source.get(field, "") or "")
    for row in nested.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, col_width in zip(row.cells, widths):
            cell.width = Inches(col_width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    funding_cell.paragraphs[-1].add_run().add_break()  # Blank line after the example.
    return nested


def generate_summary_docx(metadata, proposal_data, cicadas, *, allow_label_placeholders=False):
    doc = Document()

    nickname = ""
    if metadata['Dataset']:
        nickname = metadata['Dataset'][0].get('dataset_short_name', '')
        long_name = metadata['Dataset'][0].get('dataset_long_name', '')

    doc.add_heading(f"{nickname or 'Dataset submission summary'} (private)", level=1)
    doc.add_paragraph(
        'Public page URL: <insert here>\n'
        'Proposal type: <New Collection or New Analysis>\n'
        'Fully Executed DSA: <link to PDF>'
    )

    # [H1] Basic information:
    doc.add_heading('Basic information', level=1)
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'

    def add_row(t, label, value):
        table_row = t.add_row()
        table_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        row = table_row.cells
        row[0].text = label
        row[0].paragraphs[0].runs[0].bold = True
        row[1].text = str(value) if value is not None else ""
        return row

    # POC names, phones, emails
    prop_data = proposal_data
    poc_info = []
    for role in ["Scientific", "Technical", "Legal"]:
        name = prop_data.get(f"{role} POC Name", "")
        email = prop_data.get(f"{role} POC Email", "")
        phone = str(prop_data.get(f"{role} POC Phone", ""))
        if name:
            line = f"{role}: {name} ({email})"
            if phone and phone.lower() != "nan" and phone.strip():
                line += f" {phone}"
            poc_info.append(line.strip())

    add_row(table, "POC Information", "\n\n".join(poc_info))
    add_row(table, "Submission schedule and deadlines", prop_data.get('Time Constraints', ''))
    add_row(table, "Expected patient count", prop_data.get('number_of_subjects', ''))
    add_row(table, "Expected disk size", prop_data.get('disk_space', ''))

    # Expected types of data
    types = []
    file_formats = prop_data.get('file_formats', '')

    # Map formats to types for better display
    format_map = {}
    if file_formats:
        for part in str(file_formats).split(';'):
            if ' - ' in part:
                t, f = part.split(' - ', 1)
                format_map[t.strip()] = f.strip()

    def get_formatted_type(label, key):
        val = prop_data.get(key)
        if not val: return None

        items = []
        if isinstance(val, list):
            items = val
        elif isinstance(val, str):
            if val.startswith('[') and val.endswith(']'):
                try:
                    items = ast.literal_eval(val)
                except:
                    items = [val]
            else:
                items = [v.strip() for v in val.split(',')]

        formatted_items = []
        for item in items:
            if item in format_map:
                formatted_items.append(f"{item} ({format_map[item]})")
            else:
                formatted_items.append(item)
        return f"{label}: {', '.join(formatted_items)}"

    img_type = get_formatted_type("Images", "image_types")
    if img_type: types.append(img_type)

    supp_type = get_formatted_type("Supporting", "supporting_data")
    if supp_type: types.append(supp_type)

    der_type = get_formatted_type("Derived", "derived_types")
    if der_type: types.append(der_type)

    from tcia_proposals.proposal_metadata import parse_rows
    component_rows = parse_rows(prop_data.get('data_components'))
    add_row(table, "Expected types of data", tracking_component_text(component_rows) if component_rows else "\n".join(types))
    add_row(table, "Approximate date range of study execution", "")

    # [H1] DataCite Information:
    doc.add_heading('DataCite Information', level=1)
    table_dc = doc.add_table(rows=0, cols=2)
    table_dc.style = 'Table Grid'

    # Authors (Last Name, First Name (ORCiD ID if provided) for each author in the correct order)
    invs = sorted(metadata.get('Investigator', []), key=lambda x: int(x.get('author_order', 999)) if str(x.get('author_order', '')).isdigit() else 999)
    author_strings = []
    for inv in invs:
        auth_str = f"{inv.get('last_name', '')}, {inv.get('first_name', '')}"
        if inv.get('person_orcid'):
            auth_str += f" ({inv.get('person_orcid')})"
        author_strings.append(auth_str)
    raw_authors = prop_data.get('Authors', '')
    if not author_strings and isinstance(raw_authors, str) and raw_authors.strip():
        author_strings = [raw_authors]
    add_row(table_dc, "Authors", "\n".join(author_strings))
    rights_placeholder(add_row(table_dc, "Rights", "")[1])

    # [H1] Wordpress Page:
    doc.add_heading('Wordpress Page', level=1)
    doc.add_heading('Dataset Information', level=2)
    table_wp2 = doc.add_table(rows=0, cols=2)
    table_wp2.style = 'Table Grid'
    add_row(table_wp2, "DOI", "Do not include https://doi.org From Datacite.  10.7937/")
    add_row(table_wp2, "Status", "Set to Ongoing or Complete")
    add_row(table_wp2, "Title", metadata['Dataset'][0].get('dataset_long_name', '') if metadata['Dataset'] else "")
    add_row(table_wp2, "Short Title", nickname)
    add_row(table_wp2, "Featured Image", "Sample Image/Figure: Used as the \"Featured Image\" in Wordpress")
    ds_meta = metadata['Dataset'][0] if metadata['Dataset'] else {}
    abstract = ds_meta.get('dataset_abstract') or prop_data.get('Abstract', '')
    add_row(table_wp2, "Abstract", abstract)
    add_row(table_wp2, "CICADAS Description", "Insert link to google doc (draft where we iterate with the submitter), then populate with the final description when it's ready.")

    add_row(table_wp2, "Acknowledgements", ds_meta.get('acknowledgements', ''))

    # Structured curation sources can be extended beyond the two blank columns.
    funding_sources = parse_rows(ds_meta.get('funding_sources'))
    if not funding_sources and (ds_meta.get('funding_agency') or ds_meta.get('grant_id')):
        funding_sources = [{'funderName': ds_meta.get('funding_agency', ''),
                            'awardNumber': ds_meta.get('grant_id', '')}]
    add_funding_row(table_wp2, doc, funding_sources)

    prog_meta = metadata['Program'][0] if metadata['Program'] else {}
    add_row(table_wp2, "Program", prog_meta.get('program_short_name') or 'Community')

    doc.add_heading('Details', level=2)
    table_wp3 = doc.add_table(rows=0, cols=2)
    table_wp3.style = 'Table Grid'
    def labels_for_summary(source, field):
        try:
            return [review_term(label) for label in exact_dataset_labels(prop_data.get(source), field)]
        except ValueError:
            if not allow_label_placeholders:
                raise
            # Imported CRDC vocabularies can differ from TCIA. Leave a review
            # placeholder instead of assigning a label or blocking draft outputs.
            return ['<TCIA label review needed; see original proposal>']

    diseases = labels_for_summary('diagnosis', 'disease_type')
    locations = labels_for_summary('disease_site', 'disease_location')
    species = labels_for_summary('species', 'species')
    add_row(table_wp3, "Cancer Types", "\n".join(diseases))
    add_row(table_wp3, "Cancer Locations", "\n".join(locations))
    add_row(table_wp3, "Species", ", ".join(species) if species else "")
    add_row(table_wp3, "Subjects", prop_data.get('number_of_subjects', ''))

    doc.add_heading('Data Access', level=2)
    table_wp4 = doc.add_table(rows=0, cols=2)
    table_wp4.style = 'Table Grid'
    add_row(table_wp4, "Downloads", tracking_component_text(component_rows) if component_rows else '')
    add_row(table_wp4, "Make New Version", "SAVE YOUR CHANGES to this post before clicking this button or you will lose any changes by clicking the Update button (upper-right).  Clones the current version to the Previous Versions tables, and updates Version Number and Date Updated.")
    add_row(table_wp4, "Version Number", "The version number updates programmatically with the Make New Version button.")
    add_row(table_wp4, "Date Updated", "Updates programmatically with the Make New Version button, but can be changed manually.")
    add_row(table_wp4, "Version Change Log", "Before changing the version log, make a new version using Make New Version button, above.")

    doc.add_heading('Data Access Supplemental', level=2)
    table_wp5 = doc.add_table(rows=0, cols=2)
    table_wp5.style = 'Table Grid'
    add_row(table_wp5, "Download Info", "")
    external_resources = parse_rows(prop_data.get('external_resources'))
    resource_cell = add_row(table_wp5, "External Resources", "")[1]
    for index, resource in enumerate(external_resources):
        paragraph = resource_cell.paragraphs[0] if index == 0 else resource_cell.add_paragraph()
        paragraph.style = 'List Bullet'
        context = str(resource.get('context') or '').strip()
        paragraph.add_run(resource['url'] + (f" -- {context}" if context else ''))
    add_row(table_wp5, "External Resource Data Types", "\n".join(
        f"{i}. {resource['category']}"
        + (f" ({resource['other_category']})" if resource.get('other_category') else '')
        for i, resource in enumerate(external_resources, 1)))

    # Related Datasets (Move 'Dataset' type related works here)
    all_rel_works = metadata.get('Related_Work', [])
    rel_datasets = [w for w in all_rel_works if w.get('publication_type') == 'Dataset']
    rd_text = "\n".join([f"{d.get('title')} (DOI: {d.get('DOI')})" for d in rel_datasets])
    if not rd_text:
        raw_sources = prop_data.get('collections_analyzed', '')
        rd_text = raw_sources if isinstance(raw_sources, str) else ''
    add_row(table_wp5, "Related Datasets", rd_text)

    doc.add_heading('Citations and Data Usage Policy', level=2)
    table_wp6 = doc.add_table(rows=0, cols=2)
    table_wp6.style = 'Table Grid'

    # Citations (All other related works)
    other_rel_works = [w for w in all_rel_works if w.get('publication_type') != 'Dataset']
    citation_text = ""
    if other_rel_works:
        for w in other_rel_works:
            citation_text += f"{w.get('authorship', 'Unknown Authors')} ({w.get('year_of_publication', 'n.d.')}). {w.get('title', 'No Title')}. {w.get('journal_citation', '')} DOI: {w.get('DOI', 'No DOI')}\n\n"
    else:
        citation_text = "Source: Default Crosscite output using DOI\n\nThis is a pop out that you add:\n1. Dataset citation\n2. Data descriptor if available\n3. Any required acknowledgement"

    add_row(table_wp6, "Citations", citation_text.strip())

    # [H1] Issue Tracking
    doc.add_heading('Issue Tracking', level=1)
    table_it = doc.add_table(rows=0, cols=2)
    table_it.style = 'Table Grid'
    add_row(table_it, "Date Approved by CIP", "")
    add_row(table_it, "Kickoff email date", "")

    docx_buf = BytesIO()
    add_reference_footer(doc, proposal_data)
    doc.save(docx_buf)
    docx_buf.seek(0)
    return docx_buf
