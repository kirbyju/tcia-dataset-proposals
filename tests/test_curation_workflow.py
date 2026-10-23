from io import BytesIO
import unittest
import zipfile

import pandas as pd
from docx import Document

from tcia_proposals.curation_workflow import (
    build_model_package, import_proposal, investigator_warnings, metadata_gaps,
    normalize_investigators, parse_authors,
)
from tcia_proposals.mdf_parser import get_mdf_resources
from tcia_proposals.summary_document import generate_summary_docx


class CurationWorkflowTests(unittest.TestCase):
    def test_author_parser_preserves_names_and_cleans_orcid_urls(self):
        rows = parse_authors('(de la Cruz, Maria Isabel) - https://orcid.org/0000-0002-1825-0097; Jean Luc Picard\nConsortium')
        self.assertEqual(rows[0]['first_name'], 'Maria Isabel')
        self.assertEqual(rows[0]['last_name'], 'de la Cruz')
        self.assertEqual(rows[0]['person_orcid'], '0000-0002-1825-0097')
        self.assertEqual([r['author_order'] for r in rows], [1, 2, 3])
        self.assertEqual(len(investigator_warnings(rows)), 1)
        self.assertEqual(parse_authors('Picard, Jean Luc', 'Family, Given')[0]['first_name'], 'Jean Luc')

    def test_author_order_and_identity_warnings(self):
        rows = normalize_investigators([
            {'author_order': 3, 'first_name': 'Last', 'last_name': 'Example', 'person_orcid': '0000-0002-1825-0097'},
            {'author_order': 1, 'first_name': 'First', 'last_name': 'Example', 'person_orcid': '0000-0002-1825-0097'},
            {'author_order': 5, 'first_name': 'Bad', 'last_name': 'ORCID', 'person_orcid': '0000-0000-0000-0002'},
            {'author_order': 6, 'first_name': float('nan')},
        ])
        self.assertEqual([r['first_name'] for r in rows], ['First', 'Last', 'Bad'])
        self.assertTrue(any('repeated ORCID' in w for w in investigator_warnings(rows)))
        self.assertTrue(any('checksum' in w for w in investigator_warnings(rows)))

    def test_isolated_package_uses_actual_ids_and_does_not_mutate_import(self):
        schema, _, relationships = get_mdf_resources('tcia_proposals/resources')
        metadata = {'Program': [{'program_id': 'program-1'}],
                    'Dataset': [{'dataset_id': 'dataset-1', 'dataset_short_name': 'Example'}],
                    'Investigator': [{'author_order': 1, 'first_name': 'Ada', 'last_name': 'Example'}],
                    'Related_Work': []}
        with zipfile.ZipFile(BytesIO(build_model_package(metadata, schema, relationships, 'Example'))) as package:
            ds = pd.read_csv(package.open('Example_dataset.tsv'), sep='\t')
            inv = pd.read_csv(package.open('Example_investigator.tsv'), sep='\t')
            self.assertEqual(ds['program.program_id'][0], 'program-1')
            self.assertEqual(inv['dataset.dataset_id'][0], 'dataset-1')
            self.assertFalse(any(name.endswith('.docx') for name in package.namelist()))
        self.assertNotIn('dataset.dataset_id', metadata['Investigator'][0])

    def test_unmapped_crdc_terms_warn_and_leave_document_placeholders(self):
        proposal = {'Title': 'Example', 'Nickname': 'Example', 'diagnosis': ['A source vocabulary term']}
        metadata = {'Program': [], 'Dataset': [{'dataset_short_name': 'Example'}], 'Investigator': [], 'Related_Work': []}
        warnings = metadata_gaps(proposal, metadata, {}, 'source.xlsx')
        self.assertTrue(any('A source vocabulary term' in w for w in warnings))
        self.assertTrue(any('CRDC-only' in w for w in warnings))
        doc = Document(generate_summary_docx(metadata, proposal, {}, allow_label_placeholders=True))
        row = next(r for table in doc.tables for r in table.rows if r.cells[0].text == 'Cancer Types')
        self.assertIn('TCIA label review needed', row.cells[1].text)
        self.assertEqual(proposal['diagnosis'], ['A source vocabulary term'])

    def test_wrong_tsv_is_not_imported_as_a_proposal(self):
        with self.assertRaisesRegex(ValueError, 'No proposal title'):
            import_proposal(b'first_name\tlast_name\nAda\tExample\n', 'investigators.tsv', {})
