"""Exercise the shared validator summary generator without starting Streamlit."""
from io import BytesIO
from pathlib import Path
import unittest

from docx import Document

ROOT = Path(__file__).resolve().parents[1]


class Session(dict):
    __getattr__ = dict.__getitem__


def generate_summary(short_title="Example-Collection", proposal=None, dataset=None, program=None):
    state = Session(
        metadata={"Dataset": [{"dataset_short_name": short_title,
                                "dataset_long_name": "Example imaging dataset", **(dataset or {})}],
                  "Investigator": [], "Program": program or [], "Related_Work": []},
        proposal_raw_data=proposal or {}, cicadas={},
    )
    from tcia_proposals.summary_document import generate_summary_docx
    return generate_summary_docx(state.metadata, state.proposal_raw_data, state.cicadas)


class ValidatorSummaryTests(unittest.TestCase):
    def row(self, doc, name):
        return next(row for table in doc.tables for row in table.rows if row.cells[0].text == name)

    def test_content_from_structured_proposal_and_new_placeholders(self):
        import json
        proposal = {
            "Abstract": "Abstract from the proposal.",
            "diagnosis": "['Malignant Lung Neoplasm', 'Other - Rare custom diagnosis']",
            "disease_site": '["Lung", "Other - Custom site"]', "species": '["Human", "Mouse"]',
            "data_components": json.dumps([
                {"category": "Radiology Images", "data_types": ["CT", "MR"], "file_format": "DICOM"},
                {"category": "Radiology Images", "data_types": ["MR"], "file_format": "NIfTI"},
            ]),
            "external_resources": json.dumps([
                {"category": "Genomics", "url": "https://example.org/genomics", "context": "Matched sequencing data"},
                {"category": "Software/Source Code", "url": "https://example.org/code"},
            ]),
        }
        doc = Document(generate_summary(proposal=proposal))
        self.assertEqual(self.row(doc, "Abstract").cells[1].text, proposal["Abstract"])
        self.assertTrue(self.row(doc, "CICADAS Description").cells[1].text.startswith("Insert link to google doc"))
        self.assertEqual(self.row(doc, "Program").cells[1].text, "Community")
        self.assertEqual(self.row(doc, "Cancer Types").cells[1].text,
                         "Malignant Lung Neoplasm\nOther -- Requires FNL Review -- Rare custom diagnosis")
        self.assertEqual(self.row(doc, "Cancer Locations").cells[1].text, "Lung\nOther -- Requires FNL Review -- Custom site")
        self.assertEqual(self.row(doc, "Species").cells[1].text, "Human, Mouse")
        downloads = self.row(doc, "Downloads").cells[1].text
        self.assertIn("Category: Radiology Images\nType(s): CT, MR\nFormat: DICOM", downloads)
        self.assertIn("Category: Radiology Images\nType(s): MR\nFormat: NIfTI", downloads)
        self.assertNotIn("to assign/confirm labels.", downloads)
        self.assertEqual(self.row(doc, "Download Info").cells[1].text, "")
        self.assertEqual(self.row(doc, "External Resources").cells[1].text,
                         "https://example.org/genomics -- Matched sequencing data\nhttps://example.org/code")
        self.assertTrue(all(p.style.name == "List Bullet" for p in self.row(doc, "External Resources").cells[1].paragraphs))
        self.assertEqual(self.row(doc, "External Resource Data Types").cells[1].text,
                         "1. Genomics\n2. Software/Source Code")
        headings = [p.text for p in doc.paragraphs]
        row_names = [r.cells[0].text for t in doc.tables for r in t.rows]
        self.assertNotIn("Add New Collection", headings)
        for removed in ("New Collection Title", "Summary", "Supporting Data", "Additional Resources"):
            self.assertNotIn(removed, row_names)

    def test_component_blocks_other_terms_and_contact_spacing(self):
        import json
        rows = [
            {"category": "Radiology Images", "data_types": ["RTIMAGE", "MR", "CT"], "file_format": "DICOM"},
            {"category": "Clinical Data", "data_types": ["Protocol", "Treatment", "Diagnosis"], "file_format": "TSV"},
            {"category": "Other", "other_category": "Proposed category", "data_types": ["Other"],
             "other_data_types": "Proposed type", "file_format": "Other", "other_file_format": "Proposed format"},
        ]
        proposal = {"data_components": json.dumps(rows), **{
            f"{role} POC Name": f"{role} contact" for role in ("Scientific", "Technical", "Legal")}}
        doc = Document(generate_summary(proposal=proposal))
        expected = ("Category: Radiology Images\nType(s): RTIMAGE, MR, CT\nFormat: DICOM\n\n"
                    "Category: Clinical Data\nType(s): Protocol, Treatment, Diagnosis\nFormat: TSV\n\n"
                    "Category: Other -- Requires FNL Review -- Proposed category\n"
                    "Type: Other -- Requires FNL Review -- Proposed type\n"
                    "Format: Other -- Requires FNL Review -- Proposed format")
        for name in ("Expected types of data", "Downloads"):
            self.assertEqual(self.row(doc, name).cells[1].text, expected)
        poc = self.row(doc, "POC Information").cells[1].text
        self.assertEqual(poc.count('\n\n'), 2)
        funding_cell = next(r.cells[0] for t in doc.tables for r in t.rows if r.cells[0].tables)
        self.assertEqual(funding_cell.paragraphs[1].text, '')
        self.assertEqual(funding_cell.paragraphs[-1].text, '\n')

    def test_rights_links_and_blank_structured_funding(self):
        from tcia_proposals.proposal_metadata import FUNDING_PROPERTIES
        doc = Document(generate_summary())
        targets = {rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith('/hyperlink')}
        self.assertIn("https://creativecommons.org/licenses/by/4.0/", targets)
        self.assertIn("https://www.cancerimagingarchive.net/nih-controlled-data-access-policy/", targets)
        funding = next(row.cells[0].tables[0] for table in doc.tables for row in table.rows if row.cells[0].tables)
        self.assertEqual([c.text for c in funding.rows[0].cells],
                         ["Datacite Property", "Funding Source 1", "Funding Source 2"])
        self.assertEqual([r.cells[0].text for r in funding.rows[1:]], list(FUNDING_PROPERTIES))
        self.assertTrue(all(not c.text for r in funding.rows[1:] for c in r.cells[1:]))
        self.assertEqual(self.row(doc, "Acknowledgements").cells[1].text, "")

    def test_label_mismatches_are_flagged_without_fuzzy_mapping(self):
        from tcia_proposals.summary_document import exact_dataset_labels
        self.assertEqual(exact_dataset_labels("[]", "species"), [])
        for field, label in (("disease_type", "lung cancer"), ("species", "Other - custom"),
                             ("disease_location", "lungs")):
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, "Unrecognized"):
                exact_dataset_labels(label, field)
        with self.assertRaisesRegex(ValueError, "Unrecognized"):
            generate_summary(proposal={"diagnosis": '["lung cancer"]'})

    def test_existing_program_and_curation_abstract_are_preserved(self):
        doc = Document(generate_summary(proposal={"Abstract": "Original abstract"},
                                       dataset={"dataset_abstract": "Curated abstract"},
                                       program=[{"program_short_name": "TCGA"}]))
        self.assertEqual(self.row(doc, "Abstract").cells[1].text, "Curated abstract")
        self.assertEqual(self.row(doc, "Program").cells[1].text, "TCGA")

    def test_direct_summary_keeps_raw_authors_and_analysis_source_names(self):
        doc = Document(generate_summary(proposal={
            "Authors": "Example, Alice - 0000-0000-0000-0001; Example, Bob",
            "collections_analyzed": "Example source Collection",
        }))
        self.assertEqual(self.row(doc, "Authors").cells[1].text,
                         "Example, Alice - 0000-0000-0000-0001; Example, Bob")
        self.assertEqual(self.row(doc, "Related Datasets").cells[1].text, "Example source Collection")

    def test_short_title_is_first_h1_followed_by_placeholders(self):
        doc = Document(generate_summary())
        self.assertEqual(doc.paragraphs[0].text, "Example-Collection (private)")
        self.assertEqual(doc.paragraphs[0].style.name, "Heading 1")
        self.assertEqual(doc.paragraphs[1].text.splitlines(), [
            "Public page URL: <insert here>",
            "Proposal type: <New Collection or New Analysis>",
            "Fully Executed DSA: <link to PDF>",
        ])
        self.assertEqual(doc.paragraphs[2].text, "Basic information")

    def test_cip_date_precedes_kickoff_with_blank_values(self):
        doc = Document(generate_summary())
        self.assertEqual([[c.text for c in row.cells] for row in doc.tables[-1].rows], [
            ["Date Approved by CIP", ""], ["Kickoff email date", ""],
        ])

    def test_unset_short_title_has_readable_fallback(self):
        self.assertEqual(Document(generate_summary("")).paragraphs[0].text,
                         "Dataset submission summary (private)")


if __name__ == "__main__":
    unittest.main()
