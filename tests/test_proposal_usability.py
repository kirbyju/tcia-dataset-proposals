import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch

import pandas as pd
from streamlit.testing.v1 import AppTest

from tcia_proposals.proposal_metadata import (
    category_options, data_type_options, data_type_display, load_catalog,
    load_proposal_model, normalize_resource_url, validate_components, validate_resources,
)

ROOT = Path(__file__).resolve().parents[1]


class ProposalSelectionRulesTests(unittest.TestCase):
    def test_order_and_category_type_relationships(self):
        catalog = load_catalog()
        self.assertEqual(category_options(catalog), ["Radiology Images", "Histopathology Images",
                         "Image Annotations", "Clinical Data", "Other"])
        self.assertIn("CT", data_type_options(catalog, "Radiology Images"))
        self.assertNotIn("Treatment", data_type_options(catalog, "Radiology Images"))
        self.assertIn("Whole Slide Image", data_type_options(catalog, "Histopathology Images"))
        self.assertIn("SEG", data_type_options(catalog, "Image Annotations"))
        self.assertIn("Follow-Up", data_type_options(catalog, "Clinical Data"))
        self.assertEqual(data_type_display("CT"), "CT (Computed Tomography)")
        self.assertTrue(validate_components([{"category": "Clinical Data", "data_types": ["CT"], "file_format": "CSV"}]))
        self.assertEqual(json.loads((ROOT / "docs/labels/proposal-data-model.json").read_text()), load_proposal_model())

    def test_url_normalization_preserves_explicit_schemes_and_paths(self):
        self.assertEqual(normalize_resource_url(" gdc.nci.nih.gov "), "https://gdc.nci.nih.gov")
        self.assertEqual(normalize_resource_url("example.org/path?q=A#part"), "https://example.org/path?q=A#part")
        self.assertEqual(normalize_resource_url("http://example.org"), "http://example.org")
        for url in ("javascript:alert(1)", "mailto:name@example.org", "[invalid", "not a URL"):
            value = normalize_resource_url(url)
            self.assertTrue(validate_resources([{"category": "Genomics", "url": value}]))


class ProposalUsabilityTests(unittest.TestCase):
    def setUp(self):
        self.reference = patch('tcia_proposals.proposal_metadata.load_form_reference',
                               return_value={'catalog': load_catalog(), 'proposal_model': load_proposal_model()})
        self.reference.start()
        self.addCleanup(self.reference.stop)

    def app(self):
        app = AppTest.from_file(str(ROOT / "tcia-dataset-proposal.py"), default_timeout=30).run()
        self.assertFalse(app.exception)
        return app

    def test_defaults_faces_warning_details_and_category_switch(self):
        app = self.app()
        for key in ("faces", "published_elsewhere", "time_constraints"):
            self.assertEqual(app.radio(key=key).value, "No")
        self.assertNotIn("published_elsewhere_details", [w.key for w in app.text_area])
        app.radio(key="faces").set_value("Yes").run()
        self.assertTrue(any("nih-controlled-data-access-policy" in w.value for w in app.warning))
        app.radio(key="published_elsewhere").set_value("Yes").run()
        app.radio(key="time_constraints").set_value("Yes").run()
        self.assertIn("published_elsewhere_details", [w.key for w in app.text_area])
        self.assertIn("time_constraints_details", [w.key for w in app.text_area])
        app.selectbox(key="component_0_category").select("Radiology Images").run()
        self.assertIn("CT (Computed Tomography)", app.multiselect(key="component_0_types").options)
        self.assertNotIn("Treatment", app.multiselect(key="component_0_types").options)
        self.assertNotIn("ZIP", app.selectbox(key="component_0_format").options)
        app.multiselect(key="component_0_types").set_value(["CT", "Acquisition Metadata"]).run()
        app.selectbox(key="component_0_category").select("Histopathology Images").run()
        self.assertFalse(app.exception)
        self.assertEqual(app.multiselect(key="component_0_types").value, ["Acquisition Metadata"])
        self.assertTrue(any("cleared" in w.value for w in app.info))

    def test_manuscript_uploader_resets_and_keeps_added_files(self):
        first = io.BytesIO(b"fictional first manuscript")
        first.name = "first.pdf"
        second = io.BytesIO(b"fictional second manuscript")
        second.name = "second.pdf"
        files = {"ms_file_new_0": first, "ms_file_new_1": second}
        rendered_keys = []
        def uploader(label, *, key, **kwargs):
            import streamlit as st
            rendered_keys.append(key)
            st.session_state[key] = files.get(key)
            return files.get(key)
        with patch("streamlit.file_uploader", side_effect=uploader):
            app = self.app()
            app.radio(key="has_manuscripts").set_value("Yes").run()
            app.radio(key="ms_type_new").set_value("Upload File").run()
            app.text_input(key="disk_space").set_value("20 GB").run()
            next(b for b in app.button if b.label == "Add File to List").click().run()
            self.assertFalse(app.exception)
            self.assertEqual(app.session_state["ms_file_version"], 1)
            self.assertEqual(app.session_state["manuscript_list"][0]["file"].name, "first.pdf")
            next(b for b in app.button if b.label == "Add File to List").click().run()
            self.assertFalse(app.exception)
            self.assertEqual([r["file"].name for r in app.session_state["manuscript_list"]], ["first.pdf", "second.pdf"])
            self.assertEqual(rendered_keys[-1], "ms_file_new_2")
            self.assertEqual(app.text_input(key="disk_space").value, "20 GB")

    def test_submission_preserves_codes_normalized_url_and_yes_details(self):
        app = self.app()
        app.selectbox(key="component_0_category").select("Radiology Images").run()
        app.multiselect(key="component_0_types").set_value(["CT", "MR"])
        app.selectbox(key="component_0_format").select("DICOM").run()
        next(b for b in app.button if b.label == "Add external resource").click().run()
        app.selectbox(key="resource_0_category").select("Genomics")
        app.text_input(key="resource_0_url").set_value("gdc.nci.nih.gov").run()
        app.text_input(key="resource_0_context").set_value("Matched genomic profiles")
        self.assertEqual(app.text_input(key="resource_0_url").value, "https://gdc.nci.nih.gov")
        app.radio(key="published_elsewhere").set_value("Yes").run()
        app.text_area(key="published_elsewhere_details").set_value("Previously shared at an example repository.")
        for key, value in {
            "sci_poc_name": "Example Scientist", "sci_poc_email": "scientist@example.edu",
            "tech_poc_name": "Example Technician", "tech_poc_email": "technician@example.edu",
            "legal_poc_name": "Example Signatory", "legal_poc_email": "signatory@example.edu",
            "title": "Example imaging dataset", "nickname": "Example-Imaging",
            "disk_space": "20 GB", "number_of_subjects": "100",
        }.items():
            app.text_input(key=key).set_value(value)
        app.text_area(key="abstract").set_value("Fictional imaging proposal for software verification.")
        app.text_area(key="authors_raw_input").set_value("Example, Scientist")
        app.text_area(key="modifications").set_value("Identifying metadata and pixel content reviewed.")
        app.multiselect(key="diagnosis_ms").set_value(["Malignant Lung Neoplasm"])
        app.multiselect(key="disease_site_ms").set_value(["Lung"])
        app.multiselect(key="why_tcia").set_value(["To facilitate collaboration"]).run()
        next(b for b in app.button if b.label == "Generate Proposal Documents").click().run()
        self.assertFalse(app.exception)
        self.assertFalse(app.error)
        row = pd.read_csv(io.StringIO(app.session_state["proposal_files"]["tsv"]), sep="\t").iloc[0]
        self.assertEqual(json.loads(row['data_components'])[0]['data_types'], ['CT', 'MR'])
        self.assertEqual(json.loads(row['external_resources'])[0]['url'], 'https://gdc.nci.nih.gov')
        self.assertEqual(json.loads(row['external_resources'])[0]['context'], 'Matched genomic profiles')
        self.assertEqual(row['Published Elsewhere'], 'Yes - Previously shared at an example repository.')
        self.assertEqual(row['Time Constraints'], 'No')
        self.assertEqual(row['faces'], 'No')
        for removed in ('image_types', 'supporting_data', 'derived_types', 'file_formats', 'software_code', 'software_details'):
            self.assertNotIn(removed, row.index)
        from docx import Document
        doc = Document(io.BytesIO(app.session_state['proposal_files']['docx']))
        self.assertEqual(doc.paragraphs[0].text, 'TCIA Dataset Proposal: Example imaging dataset')
        body = '\n'.join(cell.text for table in doc.tables for r in table.rows for cell in r.cells)
        self.assertNotIn('Data model reference version', body)
        self.assertIn('https://gdc.nci.nih.gov -- Matched genomic profiles', body)
        footer = '\n'.join(p.text for section in doc.sections for p in section.footer.paragraphs)
        self.assertIn(row['label_catalog_version'], footer)
        self.assertIn(row['proposal_model_version'], footer)


if __name__ == "__main__":
    unittest.main()
