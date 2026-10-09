import hashlib
import io
import unittest

from openpyxl import load_workbook

from tcia_proposals.crdc_converter import (
    CRDC_TEMPLATE_PATH,
    CRDC_TEMPLATE_SHA256,
    build_crdc_workbook,
    import_crdc_workbook,
    is_crdc_workbook,
)


class CRDCConverterTests(unittest.TestCase):
    def setUp(self):
        self.proposal = {
            "Proposal Type": "New Collection Proposal",
            "Scientific POC Name": "Ada Lovelace",
            "Scientific POC Email": "ada@example.org",
            "Technical POC Name": "Grace Hopper",
            "Technical POC Email": "grace@example.org",
            "Legal POC Name": "Katherine Johnson",
            "Legal POC Email": "katherine@example.org",
            "Title": "Synthetic imaging study",
            "Nickname": "SYNTH-IMG",
            "Abstract": "A synthetic study used to test CRDC conversion.",
            "diagnosis": ["Bone, Sarcoma", "A made-up diagnosis"],
            "image_types": ["CT"],
            "supporting_data": ["Clinical", "Genomics"],
            "file_formats": "CT - DICOM; Clinical - TSV",
            "disk_space": "12 GB",
            "number_of_subjects": "42",
            "funding_agency": "National Cancer Institute (NCI)",
            "grant_id": "U01-TEST",
        }
        self.metadata = {
            "Program": [
                {
                    "program_name": "Community",
                    "program_short_name": "Community",
                    "program_full_description": "Community-contributed imaging collections",
                }
            ],
            "Dataset": [
                {
                    "dataset_long_name": "Synthetic imaging study",
                    "dataset_short_name": "SYNTH-IMG",
                    "dataset_abstract": "A synthetic study used to test CRDC conversion.",
                    "funding_agency": "National Cancer Institute (NCI)",
                    "grant_id": "U01-TEST",
                }
            ],
            "Investigator": [
                {
                    "author_order": 1,
                    "first_name": "Ada",
                    "last_name": "Lovelace",
                    "email": "ada@example.org",
                    "organization_name": "Example University",
                    "person_orcid": "0000-0001-2345-6789",
                    "role": "Principal Investigator",
                }
            ],
            "Related_Work": [
                {
                    "title": "Synthetic data descriptor",
                    "relationship_type": "IsDerivedFrom",
                    "DOI": "10.1234/example",
                }
            ],
        }

    def test_canonical_template_has_supported_crdc_contract(self):
        self.assertTrue(CRDC_TEMPLATE_PATH.is_file())
        self.assertEqual(
            hashlib.sha256(CRDC_TEMPLATE_PATH.read_bytes()).hexdigest(),
            CRDC_TEMPLATE_SHA256,
        )
        self.assertTrue(is_crdc_workbook(CRDC_TEMPLATE_PATH))

    def test_export_preserves_template_structure_and_populates_mapped_fields(self):
        output, warnings = build_crdc_workbook(
            self.proposal,
            self.metadata,
            {"abstract": self.proposal["Abstract"]},
        )
        workbook = load_workbook(io.BytesIO(output), data_only=False)
        template = load_workbook(CRDC_TEMPLATE_PATH, data_only=False)

        self.assertEqual(workbook.sheetnames, template.sheetnames)
        self.assertEqual(
            [sheet.sheet_state for sheet in workbook.worksheets],
            [sheet.sheet_state for sheet in template.worksheets],
        )
        self.assertEqual(
            [
                len(sheet.data_validations.dataValidation)
                for sheet in workbook.worksheets
            ],
            [
                len(sheet.data_validations.dataValidation)
                for sheet in template.worksheets
            ],
        )
        self.assertEqual(
            [sheet.freeze_panes for sheet in workbook.worksheets],
            [sheet.freeze_panes for sheet in template.worksheets],
        )
        self.assertEqual(
            sum(
                cell.data_type == "f"
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
            ),
            sum(
                cell.data_type == "f"
                for sheet in template.worksheets
                for row in sheet.iter_rows()
                for cell in row
            ),
        )
        self.assertEqual(workbook["Program and Study"]["E2"].value, "Synthetic imaging study")
        self.assertEqual(workbook["Data Access and Disease"]["K2"].value, 42)
        self.assertEqual(workbook["Data Access and Disease"]["F2"].value, "Bone, Sarcoma")
        self.assertEqual(
            workbook["Data Access and Disease"]["G2"].value,
            "A made-up diagnosis",
        )
        self.assertEqual(workbook["Data Types"]["R2"].value, "DICOM")
        self.assertTrue(any("access-type" in warning for warning in warnings))

        workbook.close()
        template.close()

    def test_export_import_round_trip_keeps_supported_core_fields(self):
        output, _ = build_crdc_workbook(
            self.proposal,
            self.metadata,
            {"abstract": self.proposal["Abstract"]},
        )
        imported = import_crdc_workbook(output)

        self.assertEqual(imported["proposal_data"]["Title"], self.proposal["Title"])
        self.assertEqual(imported["proposal_data"]["Nickname"], self.proposal["Nickname"])
        self.assertEqual(imported["proposal_data"]["number_of_subjects"], 42)
        self.assertIn("Bone, Sarcoma", imported["proposal_data"]["diagnosis"])
        self.assertEqual(imported["investigators"][0]["last_name"], "Lovelace")
        self.assertEqual(imported["related_works"][0]["DOI"], "10.1234/example")

    def test_non_crdc_xlsx_is_rejected_by_detector(self):
        from openpyxl import Workbook

        workbook = Workbook()
        output = io.BytesIO()
        workbook.save(output)
        self.assertFalse(is_crdc_workbook(output.getvalue()))

    def test_tcia_serialized_lists_keep_values_that_contain_commas(self):
        self.proposal["diagnosis"] = (
            "['Adrenal Gland, Adrenocortical Carcinoma', 'Bone, Sarcoma']"
        )
        output, _ = build_crdc_workbook(self.proposal, self.metadata)
        workbook = load_workbook(io.BytesIO(output), data_only=True)
        sheet = workbook["Data Access and Disease"]
        self.assertEqual(sheet["F2"].value, "Adrenal Gland, Adrenocortical Carcinoma")
        self.assertEqual(sheet["F3"].value, "Bone, Sarcoma")
        workbook.close()


if __name__ == "__main__":
    unittest.main()
