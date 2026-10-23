"""The task choice and investigator review gate all curation downloads."""
from io import BytesIO
import json
from pathlib import Path
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from openpyxl import load_workbook
import pandas as pd
from streamlit.testing.v1 import AppTest
from tcia_proposals.crdc_converter import build_crdc_workbook

ROOT = Path(__file__).resolve().parents[1]


def uploaded_proposal(name='Example', zipped=True, analysis=False):
    raw = {
        'Title': f'{name} imaging dataset', 'Nickname': name,
        'Abstract': f'{name} abstract', 'Authors': 'Example, Alice; Example, Bob',
        'Proposal Type': 'Analysis Results Proposal' if analysis else 'New Collection Proposal',
        'diagnosis': json.dumps(['Malignant Lung Neoplasm']),
        'disease_site': json.dumps(['Lung']), 'species': json.dumps(['Human']),
        'collections_analyzed': 'Example source' if analysis else '',
        'data_components': json.dumps([{'category': 'Radiology Images', 'data_types': ['CT', 'MR'], 'file_format': 'DICOM'}]),
        'external_resources': json.dumps([{'category': 'Genomics', 'url': 'https://example.org/genomics'}]),
    }
    payload = pd.DataFrame([raw]).to_csv(sep='\t', index=False).encode()
    if zipped:
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w') as package:
            package.writestr(f'{name}_proposal_summary.tsv', payload)
        payload = buffer.getvalue()
    upload = BytesIO(payload)
    upload.name = f'{name}.zip' if zipped else f'{name}.tsv'
    return upload


def click(app, label):
    next(b for b in app.button if b.label == label).click().run()
    assert not app.exception, app.exception


class ValidatorWorkflowTests(unittest.TestCase):
    def test_landing_screen_only_selects_a_task(self):
        with patch('streamlit.file_uploader') as upload, patch('streamlit.download_button') as download:
            app = AppTest.from_file(str(ROOT / 'tcia-remapper.py')).run()
            self.assertFalse(app.exception)
            self.assertEqual([b.label for b in app.button], ['Generate curation outputs', 'Remap supporting data'])
            upload.assert_not_called()
            download.assert_not_called()
            self.assertFalse(app.radio)
            click(app, 'Remap supporting data')
            self.assertEqual(app.session_state['workflow'], 'remap')
            upload.assert_called_once()
            self.assertIn('source data file', upload.call_args.args[0])

    def test_tcia_import_reviews_authors_then_downloads_three_outputs(self):
        downloads = {}
        with patch('streamlit.file_uploader', return_value=uploaded_proposal()) as uploader, \
             patch('streamlit.download_button', side_effect=lambda label, **kw: downloads.update({label: kw}) or False):
            app = AppTest.from_file(str(ROOT / 'tcia-remapper.py'), default_timeout=30).run()
            click(app, 'Generate curation outputs')
            self.assertFalse(downloads)
            self.assertEqual([r['first_name'] for r in app.session_state['investigator_draft']], ['Alice', 'Bob'])
            click(app, 'Confirm investigators and generate outputs')
            self.assertEqual(len(downloads), 3)
            self.assertFalse(app.error)
            doc = Document(BytesIO(downloads['Download TCIA curation tracking document']['data']))
            self.assertEqual(doc.paragraphs[0].text, 'Example (private)')
            text = '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
            self.assertIn('Example, Alice\nExample, Bob', text)
            self.assertIn('Category: Radiology Images\nType(s): CT, MR\nFormat: DICOM', text)
            with zipfile.ZipFile(BytesIO(downloads['Download TSV imaging model package']['data'])) as package:
                investigators = pd.read_csv(package.open('Example_investigator.tsv'), sep='\t')
                self.assertEqual(list(investigators.first_name), ['Alice', 'Bob'])
                self.assertTrue(investigators['dataset.dataset_id'].isna().all())
                ds = pd.read_csv(package.open('Example_dataset.tsv'), sep='\t')
                self.assertFalse(any(col.startswith('cicadas_') for col in ds.columns))
                self.assertEqual(len(package.namelist()), 3)
            workbook = load_workbook(BytesIO(downloads['Download CRDC proposal workbook']['data']))
            sheet = workbook['PI and Contact']
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.cell(2, headers['Principal Investigator First name']).value, 'Alice')
            downloads.clear()
            app.run()
            self.assertEqual(len(downloads), 3)
            # Changes to an investigator invalidate the confirmation until reviewed again.
            edited = pd.DataFrame(app.session_state['investigator_draft'])
            edited.loc[0, 'first_name'] = 'Updated'
            # AppTest does not expose data_editor edits; exercise its returned
            # dataframe across reruns, matching the browser widget contract.
            with patch('streamlit.data_editor', return_value=edited):
                downloads.clear()
                app.run()
                self.assertFalse(downloads)
                click(app, 'Confirm investigators and generate outputs')
                self.assertEqual(app.session_state['metadata']['Investigator'][0]['first_name'], 'Updated')
            uploader.return_value = uploaded_proposal('Second')
            downloads.clear()
            app.run()
            self.assertFalse(app.exception)
            self.assertFalse(downloads)
            self.assertIsNone(app.session_state['confirmed_investigators'])
            self.assertEqual(app.session_state['metadata']['Dataset'][0]['dataset_short_name'], 'Second')

    def test_analysis_gaps_warn_without_blocking_drafts(self):
        downloads = {}
        with patch('streamlit.file_uploader', return_value=uploaded_proposal(zipped=False, analysis=True)), \
             patch('streamlit.download_button', side_effect=lambda label, **kw: downloads.update({label: kw}) or False):
            app = AppTest.from_file(str(ROOT / 'tcia-remapper.py'), default_timeout=30).run()
            click(app, 'Generate curation outputs')
            click(app, 'Confirm investigators and generate outputs')
            self.assertEqual(len(downloads), 3)
            self.assertTrue(any('source dataset DOI' in warning.value for warning in app.warning))

    def test_crdc_notes_only_appear_on_crdc_download(self):
        crdc_clicked = False
        def download(label, **kwargs):
            return crdc_clicked and label == 'Download CRDC proposal workbook'
        with patch('streamlit.file_uploader', return_value=uploaded_proposal()), \
             patch('streamlit.download_button', side_effect=download):
            app = AppTest.from_file(str(ROOT / 'tcia-remapper.py'), default_timeout=30).run()
            click(app, 'Generate curation outputs')
            click(app, 'Confirm investigators and generate outputs')
            self.assertTrue(any('Imaging model' in w.value for w in app.warning))
            self.assertFalse(any('CRDC' in w.value for w in app.warning))
            crdc_clicked = True
            app.run()
            self.assertFalse(app.exception)
            self.assertTrue(any('CRDC-only' in w.value for w in app.warning))
            crdc_clicked = False
            app.run()
            self.assertFalse(any('CRDC' in w.value for w in app.warning))

    def test_remapper_runs_without_proposal_or_cicadas(self):
        upload = BytesIO(b'subject\nSUBJECT-001\nSUBJECT-002\n')
        upload.name = 'subjects.csv'
        downloads = {}
        with patch('streamlit.file_uploader', return_value=upload), \
             patch('streamlit.download_button', side_effect=lambda label, **kw: downloads.update({label: kw}) or False):
            app = AppTest.from_file(str(ROOT / 'tcia-remapper.py'), default_timeout=30).run()
            click(app, 'Remap supporting data')
            app.selectbox[0].select('Subject.subject_id').run()
            click(app, '✅ Confirm Mapping')
            self.assertFalse(app.error)
            self.assertFalse(app.session_state['proposal_raw_data'])
            self.assertIn('Download Subject.tsv', downloads)
            output = pd.read_csv(BytesIO(downloads['Download Subject.tsv']['data'].encode()), sep='\t')
            self.assertEqual(list(output.subject_id), ['SUBJECT-001', 'SUBJECT-002'])
            self.assertTrue(any('missing linkages' in warning.value for warning in app.warning))

    def test_crdc_import_retains_investigator_details(self):
        payload, _ = build_crdc_workbook({'Title': 'CRDC example', 'Nickname': 'CRDC'}, {
            'Investigator': [{'first_name': 'Ada', 'last_name': 'Example', 'email': 'ada@example.org',
                              'organization_name': 'Example University', 'author_order': 1}]})
        upload = BytesIO(payload); upload.name = 'example.xlsx'
        downloads = {}
        with patch('streamlit.file_uploader', return_value=upload), \
             patch('streamlit.download_button', side_effect=lambda label, **kw: downloads.update({label: kw}) or False):
            app = AppTest.from_file(str(ROOT / 'tcia-remapper.py'), default_timeout=30).run()
            click(app, 'Generate curation outputs')
            self.assertEqual(app.session_state['investigator_draft'][0]['email'], 'ada@example.org')
            click(app, 'Confirm investigators and generate outputs')
            self.assertEqual(len(downloads), 3)
            self.assertFalse(any('CRDC-only' in w.value for w in app.warning))
            self.assertFalse(app.exception)


if __name__ == '__main__':
    unittest.main()
