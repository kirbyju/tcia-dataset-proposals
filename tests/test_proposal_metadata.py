import io
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from docx import Document
from streamlit.testing.v1 import AppTest

from tcia_proposals.proposal_metadata import (
    FUNDING_PROPERTIES, add_curation_placeholders, label_options, load_catalog,
    parse_rows, validate_components, validate_resources,
)

ROOT = Path(__file__).resolve().parents[1]


class ProposalMetadataTests(unittest.TestCase):
    def test_registered_terms_only_and_reference_agrees(self):
        catalog = load_catalog()
        self.assertIn("Not Applicable", label_options(catalog, "disease_type"))
        self.assertNotIn("Phantom", label_options(catalog, "disease_location"))
        self.assertEqual(label_options({"fields": {"test": [
            {"label": "Retired", "registry_membership": "unregistered"},
            {"label": "Current", "registry_membership": "registered"},
        ]}}, "test"), ["Current"])
        import re
        page = (ROOT / "docs/labels/index.html").read_text()
        embedded = json.loads(re.search(r'id="catalog-data">(.*?)</script>', page, re.S).group(1))
        self.assertEqual(embedded, catalog)

    def test_multi_type_rows_roundtrip_without_collapsing_formats(self):
        rows = [
            {"category": "Radiology Images", "data_types": ["CT", "MR"], "file_format": "DICOM"},
            {"category": "Radiology Images", "data_types": ["CT"], "file_format": "NIfTI"},
            {"category": "Clinical Data", "data_types": ["Follow-Up", "Demographic", "Treatment"], "file_format": "CSV"},
        ]
        self.assertEqual(validate_components(rows), [])
        self.assertEqual(parse_rows(json.dumps(rows)), rows)
        self.assertTrue(validate_components([{**rows[0], "category": ["Radiology Images", "Clinical Data"]}]))
        self.assertTrue(validate_components([{**rows[0], "file_format": ["DICOM", "NIfTI"]}]))
        self.assertTrue(validate_components([{**rows[0], "data_types": []}]))

    def test_other_requires_description_and_resources_require_web_urls(self):
        row = {"category": "Other", "data_types": ["Other"], "file_format": "Other"}
        self.assertEqual(len(validate_components([row])), 3)
        row.update(other_category="Novel images", other_data_types="Novel type", other_file_format="New format")
        self.assertEqual(validate_components([row]), [])
        self.assertEqual(validate_resources([]), [])
        self.assertEqual(validate_resources([{"category": "Genomics", "url": "https://example.org/study"}]), [])
        for url in ("", "example.org", "javascript:alert(1)", "file:///private/data", "https://"):
            self.assertTrue(validate_resources([{"category": "Genomics", "url": url}]))

    def test_blank_funding_and_acknowledgements_document(self):
        doc = Document()
        add_curation_placeholders(doc)
        buf = io.BytesIO()
        doc.save(buf)
        doc = Document(buf)
        self.assertEqual([c.text for c in doc.tables[0].rows[0].cells],
                         ["DataCite Property", "Funding Source 1", "Funding Source 2"])
        self.assertEqual([r.cells[0].text for r in doc.tables[0].rows[1:]], list(FUNDING_PROPERTIES))
        self.assertTrue(all(not c.text for r in doc.tables[0].rows[1:] for c in r.cells[1:]))
        self.assertEqual(doc.tables[1].cell(0, 0).text, "")


class ProposalFormBehaviorTests(unittest.TestCase):
    def setUp(self):
        from tcia_proposals.proposal_metadata import load_proposal_model
        self.reference = patch('tcia_proposals.proposal_metadata.load_form_reference',
                               return_value={'catalog': load_catalog(), 'proposal_model': load_proposal_model()})
        self.reference.start()
        self.addCleanup(self.reference.stop)

    def app(self):
        app = AppTest.from_file(str(ROOT / "tcia-dataset-proposal.py"), default_timeout=20).run()
        self.assertFalse(app.exception)
        return app

    def test_defaults_and_human_age_visibility(self):
        app = self.app()
        self.assertEqual(app.multiselect(key="species").value, ["Human"])
        self.assertEqual(app.multiselect(key="adult_or_childhood_study").value, ["Adult"])
        app.multiselect(key="species").set_value(["Mouse"]).run()
        self.assertFalse(app.exception)
        self.assertNotIn("adult_or_childhood_study", [w.key for w in app.multiselect])
        self.assertFalse(any("Funding" in w.label or "Acknowledgements" in w.label for w in app.text_input))

    def test_repeatable_rows_keep_values_and_other_help(self):
        app = self.app()
        app.text_input(key="disk_space").set_value("20 GB")
        app.text_input(key="number_of_subjects").set_value("100")
        app.selectbox(key="component_0_category").select("Radiology Images").run()
        app.multiselect(key="component_0_types").set_value(["CT", "MR"])
        app.selectbox(key="component_0_format").select("DICOM").run()
        next(b for b in app.button if b.label == "Add data component").click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.multiselect(key="component_0_types").value, ["CT", "MR"])
        self.assertIsNone(app.selectbox(key="component_1_format").value)
        next(b for b in app.button if b.label == "Add external resource").click().run()
        self.assertEqual(app.text_input(key="disk_space").value, "20 GB")
        self.assertEqual(app.text_input(key="number_of_subjects").value, "100")
        self.assertIn("resource_0_url", [w.key for w in app.text_input])
        app.multiselect(key="diagnosis_ms").set_value(["Other"]).run()
        self.assertIn("diagnosis_ms_other", [w.key for w in app.text_input])
        next(b for b in app.button if b.label == "Generate Proposal Documents").click().run()
        self.assertTrue(any("Other selection" in w.value for w in app.markdown))


if __name__ == "__main__":
    unittest.main()
