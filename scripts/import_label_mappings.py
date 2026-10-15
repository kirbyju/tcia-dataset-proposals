#!/usr/bin/env python3
"""Import historical Sheet mapping evidence; optionally verify codes read-only.

Download the public source workbook separately. This never treats source-table
membership, an AI recommendation, or an existing ontology code as approval of
the TCIA label or its use for an individual dataset.
"""
from argparse import ArgumentParser
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
import urllib.parse
import urllib.request

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REFS = ROOT / "skills/curate-tcia-wordpress-metadata/references"
SOURCE = "https://docs.google.com/spreadsheets/d/1KnbOO85nqbx8A3m1VLWOiUhBJY_KxkI29_JwAVgKxjs/edit"


def read_candidates(workbook, inventory):
    sources = {field: defaultdict(list) for field in ("disease_type", "disease_location")}
    configurations = [
        ("Primary Diagnosis", "disease_type", 1, 0, 2, "NCIt"),
        ("Anatomic Site", "disease_location", 0, 3, 1, "UBERON"),
        ("Primary Site Mappings for legac", "disease_location", 0, 3, 2, "UBERON"),
    ]
    for sheet, field, label_index, code_index, definition_index, ontology in configurations:
        for number, row in enumerate(workbook[sheet].iter_rows(min_row=2, values_only=True), 2):
            if not row[label_index]:
                continue
            codes = re.findall(r"UBERON:\d+" if ontology == "UBERON" else r"\bC\d+\b", str(row[code_index] or ""))
            preferred = row[1] if sheet == "Primary Site Mappings for legac" else row[label_index]
            sources[field][str(row[label_index]).strip().casefold()].append({
                "ontology": ontology, "codes": codes,
                "source_preferred_label": preferred,
                "source_definition": row[definition_index] or "",
                "source_sheet": sheet, "source_row": number,
                "source_url": SOURCE,
                "match_method": "case-normalized source label or explicit legacy crosswalk",
            })
    comparisons = {}
    for number, row in enumerate(workbook["Updated AI legacy comparison"].iter_rows(min_row=2, values_only=True), 2):
        if row[1]:
            comparisons[str(row[1]).strip().casefold()] = {
                "comparison_status": row[0], "source_status": row[3],
                "strategy": row[4], "recommended_code": row[5],
                "secondary_codes": row[6], "recommended_label": row[7],
                "note": row[12], "source_sheet": "Updated AI legacy comparison",
                "source_row": number, "source_url": SOURCE,
            }
    fields = {}
    for field in sources:
        fields[field] = {}
        for row in inventory["fields"][field]:
            key = row["label"].strip().casefold()
            candidates = sources[field].get(key, [])
            comparison = comparisons.get(key) if field == "disease_location" else None
            evidence = {"semantic_review": "needs curator decision", "mapping_candidates": candidates,
                        "review_note": "Historical mapping evidence; semantic approval is pending." if candidates else
                        "No direct label match in the historical source. Ontology search and curator review are needed."}
            if comparison:
                evidence["legacy_comparison"] = comparison
                evidence["review_note"] += (
                    f" AI comparison ({comparison['comparison_status']}): recommends "
                    f"{comparison['recommended_code'] or comparison['secondary_codes'] or 'no single code'}. "
                    + str(comparison["note"] or "") + " This recommendation is not curator approval."
                )
            if candidates and not any(c["codes"] for c in candidates):
                evidence["review_note"] = "Source provides no ontology code; do not force a disease/anatomy mapping for this context."
            fields[field][row["label"]] = evidence
    return fields


def verify(pair):
    ontology, code = pair
    url = (f"https://api-evsrest.nci.nih.gov/api/v1/concept/ncit/{code}?include=summary" if ontology == "NCIt" else
           "https://www.ebi.ac.uk/ols4/api/ontologies/uberon/terms?obo_id=" + urllib.parse.quote(code))
    result = {"ontology": ontology, "code": code, "url": url}
    try:
        with urllib.request.urlopen(url, timeout=20) as response:
            body = response.read()
        data = json.loads(body)
        result["response_sha256"] = sha256(body).hexdigest()
        if ontology == "NCIt":
            if data.get("code") != code:
                raise ValueError("Returned code does not match requested code")
            result.update(preferred_label=data["name"], version=data.get("version"), active=data.get("active"),
                          concept_status=data.get("conceptStatus"),
                          definition=next((d["definition"] for d in data.get("definitions", []) if d.get("source") == "NCI"), ""))
        else:
            terms = [term for term in data.get("_embedded", {}).get("terms", []) if term.get("obo_id") == code and term.get("ontology_name") == "uberon"]
            if len(terms) != 1:
                raise ValueError("No unique defining UBERON term returned")
            term = terms[0]
            result.update(preferred_label=term["label"], active=not term.get("is_obsolete", True),
                          definition=" ".join(term.get("description", [])), iri=term["iri"])
        result["status"] = "verified active concept" if result.get("active") is True else "needs ontology-status review"
    except Exception as error:
        result.update(status="unable to verify", error=str(error))
    return result


def build(workbook_path, inventory_path, do_verify):
    inventory = json.loads(inventory_path.read_text())
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    fields = read_candidates(workbook, inventory)
    workbook.close()
    pairs = sorted({(candidate["ontology"], code) for rows in fields.values() for row in rows.values()
                    for candidate in row["mapping_candidates"] for code in candidate["codes"]})
    checks = dict(zip(pairs, ThreadPoolExecutor(max_workers=6).map(verify, pairs))) if do_verify else {}
    for rows in fields.values():
        for row in rows.values():
            for candidate in row["mapping_candidates"]:
                candidate["ontology_checks"] = [checks.get((candidate["ontology"], code),
                    {"code": code, "status": "not checked"}) for code in candidate["codes"]]
    return {"schema_version": "1.0.0", "provenance": {
        "source_url": SOURCE, "source_workbook_sha256": sha256(workbook_path.read_bytes()).hexdigest(),
        "retrieved_at_utc": datetime.now(timezone.utc).isoformat(),
        "method": "Direct labels and explicit source crosswalks only; no fuzzy matching or inferred curator approval.",
        "ontology_api_checks_performed": do_verify,
        "note": "The AI comparison tab is review evidence, not a curator approval record."}, "fields": fields}


if __name__ == "__main__":
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--inventory", type=Path, default=REFS / "wordpress-label-inventory-2026-10-09.json")
    parser.add_argument("--output", type=Path, default=REFS / "label-mappings.json")
    parser.add_argument("--verify", action="store_true", help="Read public NCI EVS and UBERON APIs for candidate codes")
    args = parser.parse_args()
    result = build(args.workbook, args.inventory, args.verify)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({field: {"labels": len(rows), "with_source_candidates": sum(bool(v["mapping_candidates"]) for v in rows.values())}
                      for field, rows in result["fields"].items()}))
